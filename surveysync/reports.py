from __future__ import annotations

import csv
import io
import math
import re
from pathlib import Path


def parse_point_ids_text(text: str) -> list[int]:
    ids=[]
    for raw in text.splitlines():
        row=raw.strip()
        if not row or row.startswith("#"):
            continue
        first=re.split(r"[,\t\s]+", row, maxsplit=1)[0].strip().strip('"')
        if re.fullmatch(r"[+-]?\d+", first):
            value=int(first)
            if value >= 0:
                ids.append(value)
    return sorted(set(ids))


def parse_point_ids_file(path: Path) -> list[int]:
    return parse_point_ids_text(Path(path).read_text(encoding="utf-8-sig", errors="replace"))


def available_ranges(used_ids: list[int], start: int | None = None, end: int | None = None, min_run: int = 1, *, include_open_ended: bool = False) -> list[dict]:
    """Return unused numeric PointID ranges without scanning every possible integer.

    The algorithm walks the sorted occupied IDs, so a mistyped very large upper
    bound cannot create a billions-of-iterations hang. When there are no usable
    PointIDs, both explicit bounds are required. ``include_open_ended`` is used by
    compatibility views that need the range above the last occupied PointID.
    """
    used = sorted({int(x) for x in used_ids if int(x) >= 0})
    min_run = max(1, int(min_run))
    if not used:
        if start is None or end is None:
            raise ValueError("Provide both a start and end value when the file contains no usable PointIDs.")
        lo, hi = int(start), int(end)
        if lo > hi:
            lo, hi = hi, lo
        count = hi - lo + 1
        return [{"start": lo, "end": hi, "count": count}] if count >= min_run else []

    lo = int(start if start is not None else used[0])
    hi = int(end if end is not None else used[-1])
    if lo > hi:
        lo, hi = hi, lo

    ranges: list[dict] = []
    cursor = lo
    for occupied in used:
        if occupied < cursor:
            continue
        if occupied > hi:
            break
        if occupied > cursor:
            run_end = occupied - 1
            count = run_end - cursor + 1
            if count >= min_run:
                ranges.append({"start": cursor, "end": run_end, "count": count})
        cursor = occupied + 1
    if cursor <= hi:
        count = hi - cursor + 1
        if count >= min_run:
            ranges.append({"start": cursor, "end": hi, "count": count})

    if include_open_ended and end is None:
        open_start = max(cursor, used[-1] + 1, lo)
        ranges.append({"start": open_start, "end": None, "count": None})
    return ranges


def _series_label(value: int) -> str:
    base=(int(value)//1000)*1000
    return f"{base}-series"


def _internal_range_note(start: int, end: int, used: list[int]) -> str:
    previous=max((v for v in used if v < start), default=None)
    following=min((v for v in used if v > end), default=None)
    same_series=(start//1000)==(end//1000)
    if same_series and end % 1000 == 999:
        return f"Clean open sequence through the end of the {_series_label(start)} block."
    if following is not None and following % 1000 == 0:
        return f"Open gap immediately before the {_series_label(following)} sequence begins."
    if same_series and following is not None:
        return f"Open block between existing {_series_label(start)} points; the next occupied point is {following}."
    if previous is not None and following is not None:
        return f"Open block between occupied points {previous} and {following}."
    return "Unused numeric PointID block in the uploaded survey data."


def crew_range_recommendations(
    used_ids: list[int],
    *,
    min_capacity: int = 100,
    clean_block_size: int = 1000,
    max_internal_ranges: int | None = None,
    sort_order: str = "ascending",
) -> dict:
    """Build crew-friendly PointID allocation recommendations.

    The result intentionally differs from a raw gap dump. It prioritizes:
    1. a clean, open-ended block above every occupied PointID, rounded up to the
       next thousand-series boundary; and
    2. internal gaps large enough to be practical for assigning to a field crew.

    No recommended PointID overlaps an ID found in ``used_ids``. Smaller gaps are
    counted but omitted from the recommended table by default.
    """
    used=sorted({int(v) for v in used_ids if int(v) >= 0})
    if not used:
        raise ValueError("No numeric PointIDs were found in the selected survey file(s).")
    min_capacity=max(1,int(min_capacity))
    clean_block_size=max(10,int(clean_block_size))
    raw_internal=available_ranges(used, used[0], used[-1], 1)
    eligible=[r for r in raw_internal if int(r["count"]) >= min_capacity]
    order=str(sort_order or "ascending").strip().lower()
    if order not in {"ascending","descending","capacity"}:
        raise ValueError("Point-range sort order must be ascending, descending, or capacity.")
    if order == "capacity":
        eligible.sort(key=lambda r:(-int(r["count"]), int(r["start"])))
    elif order == "descending":
        eligible.sort(key=lambda r:(-int(r["start"]), -int(r["end"])))
    else:
        eligible.sort(key=lambda r:(int(r["start"]), int(r["end"])))
    if max_internal_ranges is not None:
        eligible=eligible[:max(0,int(max_internal_ranges))]

    max_used=used[-1]
    clean_open_start=((max_used//clean_block_size)+1)*clean_block_size
    open_item={
        "start":clean_open_start,
        "end":None,
        "range":f"{clean_open_start} and above",
        "capacity":None,
        "capacity_label":"Unlimited",
        "priority":"Safest",
        "notes":f"Safest option: starts a new {clean_block_size}-point series above every PointID currently in the file, preventing overlap with existing numbers.",
        "open_ended":True,
    }
    recommendations=[]
    for r in eligible:
        start,end,count=int(r["start"]),int(r["end"]),int(r["count"])
        recommendations.append({
            "start":start,
            "end":end,
            "range":f"{start} - {end}",
            "capacity":count,
            "capacity_label":f"{count:,} point" + ("" if count==1 else "s"),
            "priority":"Large gap" if count >= 500 else ("Recommended" if count >= 250 else "Usable"),
            "notes":_internal_range_note(start,end,used),
            "open_ended":False,
        })
    # Keep numeric order intuitive for crew assignment by default. The clean
    # open-ended block belongs after finite ranges for ascending sort, first for
    # descending/capacity views.
    if order == "ascending":
        recommendations.append(open_item)
    else:
        recommendations.insert(0,open_item)
    omitted=[r for r in raw_internal if int(r["count"]) < min_capacity]
    return {
        "used_count":len(used),
        "lowest_used":used[0],
        "highest_used":used[-1],
        "minimum_capacity":min_capacity,
        "sort_order":order,
        "recommended_count":len(recommendations),
        "internal_recommended_count":len(eligible),
        "smaller_gap_count":len(omitted),
        "recommendations":recommendations,
        "all_internal_ranges":raw_internal,
    }


def ranges_csv(ranges: list[dict]) -> str:
    out=io.StringIO(newline="")
    w=csv.DictWriter(out, fieldnames=["start","end","count"])
    w.writeheader(); w.writerows(ranges)
    return out.getvalue()


def crew_ranges_csv(summary: dict) -> str:
    out=io.StringIO(newline="")
    w=csv.writer(out,lineterminator="\r\n")
    w.writerow(["Available Point Range","Capacity","Priority","Notes"])
    for item in summary.get("recommendations",[]):
        w.writerow([item.get("range",""),item.get("capacity_label",""),item.get("priority",""),item.get("notes","")])
    return out.getvalue()


def crew_ranges_txt(summary: dict) -> str:
    """Plain-text crew allocation report with explicit From / To columns."""
    rows = list(summary.get("recommendations") or [])
    lines = [
        "SurveySync Available Point Ranges",
        f"Occupied numeric PointIDs: {int(summary.get('used_count') or 0):,}",
        f"Lowest used: {summary.get('lowest_used', '')}",
        f"Highest used: {summary.get('highest_used', '')}",
        f"Sort: {summary.get('sort_order', 'ascending')}",
        "",
        f"{'From':>12} {'To':>12} {'Capacity':>12}  Priority",
        f"{'-'*12} {'-'*12} {'-'*12}  {'-'*16}",
    ]
    for item in rows:
        start = str(item.get("start", ""))
        end = "and above" if item.get("end") is None else str(item.get("end"))
        capacity = "Unlimited" if item.get("capacity") is None else f"{int(item.get('capacity')):,}"
        lines.append(f"{start:>12} {end:>12} {capacity:>12}  {item.get('priority','')}")
    lines.extend(["", "Notes:"])
    for item in rows:
        lines.append(f"- {item.get('range','')}: {item.get('notes','')}")
    return "\n".join(lines).rstrip() + "\n"


def write_crew_ranges_xlsx(summary: dict, path: Path) -> Path:
    """Write the crew-allocation report to an Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Available Point Ranges"
    ws.append(["SurveySync Available Point Ranges"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Occupied numeric PointIDs", int(summary.get("used_count") or 0)])
    ws.append(["Lowest used", summary.get("lowest_used")])
    ws.append(["Highest used", summary.get("highest_used")])
    ws.append(["Sort", summary.get("sort_order", "ascending")])
    ws.append([])
    headers = ["From", "To", "Available Point Range", "Capacity", "Priority", "Notes"]
    ws.append(headers)
    for cell in ws[7]:
        cell.font = Font(bold=True)
    for item in summary.get("recommendations") or []:
        ws.append([
            item.get("start"),
            "and above" if item.get("end") is None else item.get("end"),
            item.get("range", ""),
            "Unlimited" if item.get("capacity") is None else item.get("capacity"),
            item.get("priority", ""),
            item.get("notes", ""),
        ])
    widths = [14, 14, 24, 16, 18, 72]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A8"
    wb.save(output)
    return output
