"""Strict delimited point/code import with explicit mapping and no dropped rows."""

from __future__ import annotations

import csv
import io
import math
from dataclasses import asdict, dataclass
from pathlib import Path

from .codes import CodeRule, suggest_role
from ..control_import_mapping import detect_control_mapping

MAX_POINTS = 100_000


@dataclass(frozen=True)
class TopoPoint:
    point_id: str
    northing: float
    easting: float
    elevation: float
    code: str
    source_row: int

    def to_dict(self) -> dict:
        return asdict(self)


def table_from_text(text: str, *, max_rows: int = MAX_POINTS + 1) -> list[list[str]]:
    if "\ufffd" in text:
        raise ValueError("The file contains unreadable characters. Export as UTF-8 CSV or TXT.")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        rows = [[v.strip() for v in row] for row in csv.reader(io.StringIO(text), dialect)]
    except csv.Error:
        # Whitespace PNEZD: retain the complete description as column five.
        rows = [line.strip().split(None, 4) for line in text.splitlines() if line.strip()]
    rows = [r for r in rows if any(r)]
    if not rows:
        raise ValueError("The file has no data rows.")
    if len(rows) > max_rows:
        raise ValueError(f"File exceeds {max_rows - 1:,} records. Split by acquisition session.")
    width = len(rows[0])
    if width < 2 or any(len(r) != width for r in rows):
        raise ValueError("Rows have inconsistent column counts; fix or re-export the source.")
    return rows


def survey_preview(text: str, has_header: bool | None = None) -> dict:
    rows = table_from_text(text)
    if has_header is None:
        known = {
            "point_id",
            "pointid",
            "point",
            "point_number",
            "northing",
            "easting",
            "elevation",
            "code",
            "description",
        }
        has_header = len({v.strip().lower() for v in rows[0]} & known) >= 2
    headers = rows[0] if has_header else [f"Column {i + 1}" for i in range(len(rows[0]))]
    if len(set(headers)) != len(headers) or any(not h for h in headers):
        raise ValueError("Column headers must be nonempty and unique.")
    data = rows[1:] if has_header else rows
    if not data or len(data) > MAX_POINTS:
        raise ValueError(f"Expected 1–{MAX_POINTS:,} survey records.")
    records = [dict(zip(headers, row)) for row in data]
    mapping = detect_control_mapping(headers, records[:30]).get("mapping", {})
    # Control import aliases do not include every common standalone PointID header.
    for header in headers:
        if header.lower().replace("_", "").replace(" ", "") in {
            "pointid",
            "pointnumber",
            "point",
            "pid",
        }:
            mapping.setdefault("point_id", header)
            break
    if not has_header and len(headers) >= 5:
        mapping = dict(
            zip(("point_id", "northing", "easting", "elevation", "description"), headers[:5])
        )
    return {
        "headers": headers,
        "mapping": mapping,
        "has_header": has_header,
        "row_count": len(records),
        "preview": records[:8],
        "rows": records,
    }


def parse_points(scan: dict, mapping: dict[str, str]) -> list[TopoPoint]:
    required = ("point_id", "northing", "easting", "elevation", "description")
    if any(mapping.get(key) not in scan["headers"] for key in required):
        raise ValueError("Map Point ID, Northing, Easting, Elevation and Code before analysis.")
    if len({mapping[k] for k in required}) != len(required):
        raise ValueError("Each required field must use a different source column.")
    seen: set[str] = set()
    points: list[TopoPoint] = []
    for index, row in enumerate(scan["rows"]):
        row_no = index + (2 if scan["has_header"] else 1)
        pid = row[mapping["point_id"]].strip()
        if not pid or pid in seen:
            raise ValueError(
                f"Row {row_no}: empty or duplicate PointID {pid!r}. No rows were discarded."
            )
        seen.add(pid)
        try:
            n, e, z = (float(row[mapping[k]]) for k in ("northing", "easting", "elevation"))
        except (ValueError, TypeError) as exc:
            raise ValueError(f"Row {row_no} ({pid}): invalid N/E/Z value.") from exc
        if not all(math.isfinite(v) for v in (n, e, z)):
            raise ValueError(f"Row {row_no} ({pid}): non-finite coordinate/elevation.")
        points.append(TopoPoint(pid, n, e, z, row[mapping["description"]], row_no))
    return points


def parse_code_file(path: Path) -> list[CodeRule]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [
                [str(v or "").strip() for v in row]
                for row in book.active.iter_rows(values_only=True)
            ]
        finally:
            book.close()
        rows = [r for r in rows if any(r)]
    else:
        rows = table_from_text(path.read_text(encoding="utf-8-sig"), max_rows=5001)
    if not rows or len(rows) > 5001:
        raise ValueError("Code list must contain 1–5,000 codes.")
    headers = [v.lower() for v in rows[0]]
    if "code" not in headers or "description" not in headers:
        raise ValueError("Code list requires headers: code, description, and optional role.")
    ci, di = headers.index("code"), headers.index("description")
    ri = headers.index("role") if "role" in headers else None
    result: list[CodeRule] = []
    seen: set[str] = set()
    for row in rows[1:]:
        code, description = row[ci].upper(), row[di]
        if not code or code in seen:
            raise ValueError(f"Empty or duplicate code {code!r}.")
        seen.add(code)
        role = row[ri].lower() if ri is not None and row[ri] else suggest_role(description)
        result.append(
            CodeRule.model_validate({"code": code, "description": description, "role": role})
        )
    return result
