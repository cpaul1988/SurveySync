"""ReportSync Excel Template Mapper.

Templates are preserved as immutable project sources. Mappings remain separate from
the source workbook, and rendering always writes a new workbook.
"""

from __future__ import annotations

import copy
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, column_index_from_string
from openpyxl.cell.cell import MergedCell
from zipfile import BadZipFile

from .audit import utc_now
from .project import SurveyProject, safe_name
from .reporting import register_deliverable

PLACEHOLDER = re.compile(r"{{\s*([A-Za-z0-9_.-]+)\s*}}")

ALLOWED_POINT_FIELDS = {
    "point_id",
    "northing",
    "easting",
    "elevation",
    "description",
    "point_class",
    "review_state",
    "crs",
    "horizontal_units",
    "vertical_units",
}


class ReportTemplateError(ValueError):
    pass


def _mapping_path(project: SurveyProject) -> Path:
    return project.paths.root / ".surveysync" / "report_template_mappings.json"


def _load_store(project: SurveyProject) -> dict[str, Any]:
    path = _mapping_path(project)
    if not path.is_file():
        return {"version": 1, "templates": []}
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ReportTemplateError(f"Template mapping store could not be read: {exc}") from exc
    if not isinstance(raw, dict) or not isinstance(raw.get("templates", []), list):
        raise ReportTemplateError("Template mapping store is invalid.")
    return {"version": int(raw.get("version") or 1), "templates": list(raw["templates"])}


def _save_store(project: SurveyProject, store: dict[str, Any]) -> None:
    path = _mapping_path(project)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(store, indent=2, sort_keys=True), encoding="utf-8")
    temp.replace(path)


def _immutable_source_path(project: SurveyProject, source: dict[str, Any]) -> Path:
    raw = Path(str(source.get("stored_path") or ""))
    path = raw if raw.is_absolute() else project.paths.root / raw
    if not path.is_file():
        raise ReportTemplateError("Stored template source is missing.")
    return path.resolve()


def _context(project: SurveyProject) -> dict[str, Any]:
    with project.db.connect() as conn:
        point_count = int(conn.execute("SELECT COUNT(*) FROM canonical_points").fetchone()[0])
        source_count = int(conn.execute("SELECT COUNT(*) FROM source_registry").fetchone()[0])
        open_qa = int(
            conn.execute("SELECT COUNT(*) FROM qa_issues WHERE status='OPEN'").fetchone()[0]
        )
    audit = project.db.verify_audit_chain()
    now = datetime.now(timezone.utc)
    return {
        "project": {
            "id": project.manifest.get("project_id", ""),
            "name": project.manifest.get("name", ""),
            "number": project.manifest.get("project_number", ""),
            "client": project.manifest.get("client", ""),
            "crs": project.manifest.get("crs", ""),
            "horizontal_units": project.manifest.get("horizontal_units", ""),
            "vertical_units": project.manifest.get("vertical_units", ""),
            "revision": project.manifest.get("revision", 0),
        },
        "generated": {
            "date": now.strftime("%Y-%m-%d"),
            "utc": now.isoformat(),
            "product": "SurveySync",
        },
        "counts": {
            "points": point_count,
            "sources": source_count,
            "open_qa": open_qa,
        },
        "audit": {
            "verified": bool(audit.get("ok")),
            "head_hash": audit.get("head_hash", ""),
            "event_count": audit.get("event_count", 0),
        },
    }


def _resolve(context: dict[str, Any], field: str) -> Any:
    current: Any = context
    for token in str(field or "").split("."):
        if not token:
            raise ReportTemplateError("Template field path is blank.")
        if not isinstance(current, dict) or token not in current:
            raise ReportTemplateError(f"Unknown template field: {field}")
        current = current[token]
    return current


def _scan_placeholders(path: Path) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except (OSError, ValueError, BadZipFile, KeyError) as exc:
        raise ReportTemplateError(f"Excel template could not be opened: {exc}") from exc
    found: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row * sheet.max_column > 1000000:
                raise ReportTemplateError("Template is too large for bounded inspection.")
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or "{{" not in value:
                        continue
                    fields = PLACEHOLDER.findall(value)
                    if fields:
                        found.append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "template_text": value,
                                "fields": fields,
                            }
                        )
    finally:
        workbook.close()
    return found


def inspect_excel_template(path: str | Path) -> dict[str, Any]:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise ReportTemplateError(f"Excel template was not found: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ReportTemplateError("ReportSync Template Mapper supports .xlsx and .xlsm.")
    try:
        workbook = load_workbook(
            source,
            read_only=True,
            data_only=False,
            keep_vba=source.suffix.lower() == ".xlsm",
        )
    except (OSError, ValueError, BadZipFile, KeyError) as exc:
        raise ReportTemplateError(f"Excel template could not be opened: {exc}") from exc
    try:
        sheets = [
            {
                "name": sheet.title,
                "max_row": int(sheet.max_row),
                "max_column": int(sheet.max_column),
            }
            for sheet in workbook.worksheets
        ]
        named_ranges = sorted(
            str(item.name)
            for item in workbook.defined_names.values()
            if getattr(item, "name", None)
        )
    finally:
        workbook.close()
    return {
        "path": str(source),
        "filename": source.name,
        "sheets": sheets,
        "named_ranges": named_ranges,
        "placeholders": _scan_placeholders(source),
        "supported_scalar_fields": [
            "project.id",
            "project.name",
            "project.number",
            "project.client",
            "project.crs",
            "project.horizontal_units",
            "project.vertical_units",
            "project.revision",
            "generated.date",
            "generated.utc",
            "generated.product",
            "counts.points",
            "counts.sources",
            "counts.open_qa",
            "audit.verified",
            "audit.head_hash",
            "audit.event_count",
        ],
        "supported_point_fields": sorted(ALLOWED_POINT_FIELDS),
    }


def _normalize_mapping(mapping: dict[str, Any]) -> dict[str, Any]:
    scalar_raw = mapping.get("scalar_cells") or []
    if not isinstance(scalar_raw, list):
        raise ReportTemplateError("scalar_cells must be a list.")
    scalar_cells = []
    for index, item in enumerate(scalar_raw):
        if not isinstance(item, dict):
            raise ReportTemplateError(f"Scalar mapping {index + 1} must be an object.")
        field = str(item.get("field") or "").strip()
        sheet = str(item.get("sheet") or "").strip()
        cell = str(item.get("cell") or "").strip().upper()
        if not field or not sheet or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", cell):
            raise ReportTemplateError(
                f"Scalar mapping {index + 1} requires field, sheet, and cell."
            )
        row, column = coordinate_to_tuple(cell)
        if row > 1048576 or column > 16384:
            raise ReportTemplateError("Scalar mapping exceeds Excel worksheet bounds.")
        if any(x["sheet"] == sheet and x["cell"] == cell for x in scalar_cells):
            raise ReportTemplateError("Two scalar fields cannot target the same cell.")
        scalar_cells.append({"field": field, "sheet": sheet, "cell": cell})

    table_raw = mapping.get("point_table")
    point_table = None
    if table_raw:
        if not isinstance(table_raw, dict):
            raise ReportTemplateError("point_table must be an object.")
        sheet = str(table_raw.get("sheet") or "").strip()
        start_row = int(table_raw.get("start_row") or 0)
        columns = table_raw.get("columns") or {}
        if not sheet or start_row < 1 or not isinstance(columns, dict) or not columns:
            raise ReportTemplateError("point_table requires sheet, start_row >= 1, and columns.")
        normalized_columns: dict[str, str] = {}
        for column, field in columns.items():
            col = str(column or "").strip().upper()
            point_field = str(field or "").strip()
            if not col or not re.fullmatch(r"[A-Z]{1,3}", col):
                raise ReportTemplateError(f"Invalid Excel column: {column}")
            if column_index_from_string(col) > 16384:
                raise ReportTemplateError("Point mapping exceeds Excel worksheet bounds.")
            if col in normalized_columns:
                raise ReportTemplateError("Duplicate point table column.")
            if point_field not in ALLOWED_POINT_FIELDS:
                raise ReportTemplateError(f"Unsupported point table field: {point_field}")
            normalized_columns[col] = point_field
        if start_row > 1048576:
            raise ReportTemplateError("Point table start row exceeds Excel limits.")
        point_table = {
            "sheet": sheet,
            "start_row": start_row,
            "columns": normalized_columns,
        }

    return {
        "scalar_cells": scalar_cells,
        "point_table": point_table,
        "replace_placeholders": bool(mapping.get("replace_placeholders", True)),
    }


def _source_row(project: SurveyProject, source_id: str) -> dict[str, Any]:
    with project.db.connect() as conn:
        row = conn.execute(
            "SELECT * FROM source_registry WHERE source_id=?", (source_id,)
        ).fetchone()
    if row is None:
        raise ReportTemplateError("Template source was not found in the source registry.")
    return dict(row)


def register_excel_template(
    project: SurveyProject,
    path: str | Path,
    *,
    name: str = "",
    mapping: dict[str, Any] | None = None,
) -> dict[str, Any]:
    inspection = inspect_excel_template(path)
    template_id = uuid4().hex
    learned_scalar = []
    if mapping is None:
        for placeholder in inspection["placeholders"]:
            fields = placeholder["fields"]
            if len(fields) == 1 and placeholder["template_text"].strip() == (
                "{{" + fields[0] + "}}"
            ):
                learned_scalar.append(
                    {
                        "field": fields[0],
                        "sheet": placeholder["sheet"],
                        "cell": placeholder["cell"],
                    }
                )
    normalized = _normalize_mapping(
        mapping
        or {
            "scalar_cells": learned_scalar,
            "point_table": None,
            "replace_placeholders": True,
        }
    )
    validate_mapping_targets(project, Path(path), normalized)
    source = project.import_source(
        Path(path),
        "ReportSync",
        "Original company/client Excel report template retained immutably.",
    )
    entry = {
        "template_id": template_id,
        "name": str(name or Path(path).stem).strip() or "Excel Template",
        "source_id": str(source["source_id"]),
        "original_name": inspection["filename"],
        "created_utc": utc_now(),
        "mapping": normalized,
    }
    store = _load_store(project)
    store["templates"].append(entry)
    _save_store(project, store)
    project.db.audit(
        "ReportSync",
        "REPORT_TEMPLATE_REGISTERED",
        object_type="report_template",
        object_id=template_id,
        details={
            "name": entry["name"],
            "source_id": entry["source_id"],
            "placeholder_count": len(inspection["placeholders"]),
            "learned_scalar_count": len(learned_scalar),
        },
    )
    return {**entry, "inspection": inspection}


def list_templates(project: SurveyProject) -> list[dict[str, Any]]:
    return [dict(item) for item in _load_store(project)["templates"]]


def _template_entry(project: SurveyProject, template_id: str) -> dict[str, Any]:
    for item in list_templates(project):
        if str(item.get("template_id") or "") == str(template_id):
            return item
    raise ReportTemplateError("Report template mapping was not found.")


def save_template_mapping(
    project: SurveyProject, template_id: str, mapping: dict[str, Any]
) -> dict[str, Any]:
    normalized = _normalize_mapping(mapping)
    current = _template_entry(project, template_id)
    path = _immutable_source_path(project, _source_row(project, current["source_id"]))
    validate_mapping_targets(project, path, normalized)
    store = _load_store(project)
    updated = None
    for item in store["templates"]:
        if str(item.get("template_id") or "") == str(template_id):
            item["mapping"] = normalized
            item["modified_utc"] = utc_now()
            updated = dict(item)
            break
    if updated is None:
        raise ReportTemplateError("Report template mapping was not found.")
    _save_store(project, store)
    project.db.audit(
        "ReportSync",
        "REPORT_TEMPLATE_MAPPING_SAVED",
        object_type="report_template",
        object_id=template_id,
        details={
            "scalar_count": len(normalized["scalar_cells"]),
            "point_table": normalized["point_table"],
        },
    )
    return updated


def delete_template(project: SurveyProject, template_id: str) -> dict[str, Any]:
    store = _load_store(project)
    kept = [
        item
        for item in store["templates"]
        if str(item.get("template_id") or "") != str(template_id)
    ]
    if len(kept) == len(store["templates"]):
        raise ReportTemplateError("Report template mapping was not found.")
    store["templates"] = kept
    _save_store(project, store)
    project.db.audit(
        "ReportSync",
        "REPORT_TEMPLATE_MAPPING_DELETED",
        object_type="report_template",
        object_id=template_id,
    )
    return {"template_id": template_id, "deleted": True, "source_preserved": True}


def _copy_row_style(sheet: Any, source_row: int, target_row: int, columns: list[str]) -> None:
    if target_row == source_row:
        return
    for column in columns:
        source = sheet[f"{column}{source_row}"]
        target = sheet[f"{column}{target_row}"]
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy.copy(source.alignment)
        if source.protection:
            target.protection = copy.copy(source.protection)


def _project_points(project: SurveyProject) -> list[dict[str, Any]]:
    with project.db.connect() as conn:
        return [
            dict(row)
            for row in conn.execute(
                """SELECT point_id,northing,easting,elevation,description,point_class,
                          review_state,crs,horizontal_units,vertical_units
                   FROM canonical_points ORDER BY point_id"""
            ).fetchall()
        ]


def render_excel_template(
    project: SurveyProject,
    template_id: str,
    *,
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    entry = _template_entry(project, template_id)
    source = _source_row(project, str(entry["source_id"]))
    template_path = _immutable_source_path(project, source)
    mapping = _normalize_mapping(dict(entry.get("mapping") or {}))
    keep_vba = template_path.suffix.lower() == ".xlsm"
    validate_mapping_targets(project, template_path, mapping)
    context = _context(project)
    try:
        workbook = load_workbook(template_path, data_only=False, keep_vba=keep_vba)
    except (OSError, ValueError, BadZipFile, KeyError) as exc:
        raise ReportTemplateError(f"Excel template could not be opened: {exc}") from exc

    try:
        for item in mapping["scalar_cells"]:
            if item["sheet"] not in workbook.sheetnames:
                raise ReportTemplateError(f"Worksheet not found: {item['sheet']}")
            _write_input(workbook[item["sheet"]][item["cell"]], _resolve(context, item["field"]))

        replaced_placeholders = 0
        if mapping["replace_placeholders"]:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if not isinstance(cell.value, str) or "{{" not in cell.value:
                            continue
                        original = cell.value

                        def replace(match: re.Match[str]) -> str:
                            nonlocal replaced_placeholders
                            value = _resolve(context, match.group(1))
                            replaced_placeholders += 1
                            return str(value if value is not None else "")

                        if cell.data_type == "f":
                            raise ReportTemplateError(
                                "Formula placeholders are unsafe. Map values into separate input cells instead."
                            )
                        _write_input(cell, PLACEHOLDER.sub(replace, original))

        point_count = 0
        table = mapping["point_table"]
        if table:
            if table["sheet"] not in workbook.sheetnames:
                raise ReportTemplateError(f"Worksheet not found: {table['sheet']}")
            sheet = workbook[table["sheet"]]
            points = _project_points(project)
            columns = list(table["columns"])
            start = int(table["start_row"])
            for index, point in enumerate(points):
                row_number = start + index
                _copy_row_style(sheet, start, row_number, columns)
                for column, field in table["columns"].items():
                    _write_input(sheet[f"{column}{row_number}"], point.get(field))
            if not points:
                for column in columns:
                    sheet[f"{column}{start}"] = None
            point_count = len(points)

        if output_path:
            output = Path(output_path).expanduser().resolve()
        else:
            stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
            extension = ".xlsm" if keep_vba else ".xlsx"
            output = (
                project.paths.reports
                / "TemplateExports"
                / f"{safe_name(entry['name'])}_{stamp}_{uuid4().hex[:12]}{extension}"
            )
        if output == template_path:
            raise ReportTemplateError(
                "Rendered output cannot overwrite immutable template evidence."
            )
        try:
            output.relative_to(project.paths.source.resolve())
        except ValueError:
            pass
        else:
            raise ReportTemplateError(
                "Rendered output cannot be written inside the immutable Source folder."
            )
        if output.suffix.lower() != (".xlsm" if keep_vba else ".xlsx"):
            raise ReportTemplateError("Output extension must match the template workbook type.")
        if output.exists():
            raise ReportTemplateError(
                "Output already exists. Choose a new name; original files are never overwritten."
            )
        output.parent.mkdir(parents=True, exist_ok=True)
        created = False
        try:
            with output.open("xb") as destination:
                created = True
                workbook.save(destination)
        except Exception:
            # Never delete a competing/preexisting file after an exclusive-create race.
            if created:
                output.unlink(missing_ok=True)
            raise
    finally:
        workbook.close()

    deliverable = register_deliverable(
        project,
        output,
        module="ReportSync",
        kind="excel_template_report",
        status="DRAFT",
        metadata={
            "template_id": template_id,
            "template_source_id": entry["source_id"],
            "template_source_sha256": source.get("sha256", ""),
            "mapping": mapping,
            "point_count": point_count,
            "rendered_utc": utc_now(),
        },
    )
    project.db.audit(
        "ReportSync",
        "REPORT_TEMPLATE_RENDERED",
        object_type="report_template",
        object_id=template_id,
        details={
            "output": str(output),
            "deliverable_id": deliverable["deliverable_id"],
            "status": "DRAFT",
            "point_count": point_count,
            "replaced_placeholders": replaced_placeholders,
        },
    )
    return {
        "template": entry,
        "output_path": str(output),
        "point_count": point_count,
        "replaced_placeholders": replaced_placeholders,
        "deliverable": deliverable,
    }


def _write_input(cell: Any, value: Any) -> None:
    if isinstance(cell, MergedCell) or cell.data_type == "f":
        raise ReportTemplateError(
            "Mapped inputs cannot overwrite a formula or a merged-cell continuation."
        )
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # PointIDs and client text are data, never formulas.


def validate_mapping_targets(project: SurveyProject, path: Path, mapping: dict) -> None:
    context = _context(project)
    try:
        workbook = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    except (OSError, ValueError, BadZipFile, KeyError) as exc:
        raise ReportTemplateError(f"Cannot validate template: {exc}") from exc
    try:
        targets = set()
        for item in mapping["scalar_cells"]:
            _resolve(context, item["field"])
            if item["sheet"] not in workbook.sheetnames:
                raise ReportTemplateError("Scalar mapping worksheet was not found.")
            cell = workbook[item["sheet"]][item["cell"]]
            if isinstance(cell, MergedCell) or cell.data_type == "f":
                raise ReportTemplateError(
                    "Scalar mapping targets a formula or merged-cell continuation."
                )
            targets.add((item["sheet"], item["cell"]))
        table = mapping.get("point_table")
        if table:
            if table["sheet"] not in workbook.sheetnames:
                raise ReportTemplateError("Point table worksheet was not found.")
            with project.db.connect() as conn:
                count = conn.execute("SELECT COUNT(*) FROM canonical_points").fetchone()[0]
            start = table["start_row"]
            if start + max(1, count) - 1 > 1048576 or count * len(table["columns"]) > 1000000:
                raise ReportTemplateError("Mapped point table exceeds supported workbook size.")
            for row in range(start, start + max(1, count)):
                for col in table["columns"]:
                    address = f"{col}{row}"
                    cell = workbook[table["sheet"]][address]
                    if (table["sheet"], address) in targets:
                        raise ReportTemplateError("Point table overlaps a scalar mapping.")
                    if isinstance(cell, MergedCell) or cell.data_type == "f":
                        raise ReportTemplateError(
                            "Point table would overwrite a formula or merged-cell continuation."
                        )
    finally:
        workbook.close()


def template_details(project: SurveyProject, template_id: str) -> dict:
    entry = _template_entry(project, template_id)
    path = _immutable_source_path(project, _source_row(project, entry["source_id"]))
    return {"template": entry, "inspection": inspect_excel_template(path)}
