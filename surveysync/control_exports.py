from __future__ import annotations
import csv
from pathlib import Path
from .audit import utc_now


def write_control_qc_deliverables(result: dict, reports_dir: Path) -> dict:
    from openpyxl import Workbook

    folder = Path(reports_dir)
    folder.mkdir(parents=True, exist_ok=True)
    stamp = utc_now().replace(":", "-")
    accepted_path = folder / f"Control_Accepted_{stamp}.csv"
    reshoot_path = folder / f"Control_Reshoot_{stamp}.csv"
    xlsx_path = folder / f"Control_QC_{stamp}.xlsx"
    accepted_rows = []
    reshoot_rows = []
    for item in result.get("results", []):
        selected = item.get("selected") or {}
        if item.get("status") == "PASS":
            accepted_rows.append(
                {
                    "control_id": item.get("control_id", ""),
                    "northing": selected.get("northing"),
                    "easting": selected.get("easting"),
                    "elevation": selected.get("elevation"),
                    "code": selected.get("code", ""),
                    "horizontal_residual": selected.get("max_horizontal_residual"),
                    "vertical_residual": selected.get("max_vertical_residual"),
                    "source_observations": ";".join(selected.get("point_ids") or []),
                    "numbering_flags": "; ".join(
                        (f.get("message", "") for f in item.get("numbering_flags") or [])
                    ),
                    "field_qc_status": (selected.get("field_validation") or {}).get("status", ""),
                    "minimum_time_gap_minutes": (
                        (selected.get("field_validation") or {}).get("time_separation") or {}
                    ).get("minimum_gap_minutes"),
                    "qc_status": "PASS",
                }
            )
        elif item.get("status") in {"RESHOOT", "REVIEW"}:
            reshoot_rows.append(
                {
                    "control_id": item.get("control_id", ""),
                    "reason": item.get("reason", ""),
                    "existing_shots": ";".join(item.get("existing_point_ids") or []),
                    "next_shots": ";".join(item.get("reshoot_point_ids") or []),
                    "numbering_flags": "; ".join(
                        (f.get("message", "") for f in item.get("numbering_flags") or [])
                    ),
                    "field_qc_status": (selected.get("field_validation") or {}).get("status", ""),
                    "field_failures": "; ".join(
                        (selected.get("field_validation") or {}).get("failures") or []
                    ),
                    "missing_field_metadata": "; ".join(
                        (selected.get("field_validation") or {}).get("missing_metadata") or []
                    ),
                    "horizontal_residual": selected.get("max_horizontal_residual"),
                    "vertical_residual": selected.get("max_vertical_residual"),
                }
            )
    with accepted_path.open("w", newline="", encoding="utf-8-sig") as fh:
        fields = [
            "control_id",
            "northing",
            "easting",
            "elevation",
            "code",
            "horizontal_residual",
            "vertical_residual",
            "source_observations",
            "numbering_flags",
            "field_qc_status",
            "minimum_time_gap_minutes",
            "qc_status",
        ]
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(accepted_rows)
    with reshoot_path.open("w", newline="", encoding="utf-8-sig") as fh:
        fields = [
            "control_id",
            "reason",
            "existing_shots",
            "next_shots",
            "numbering_flags",
            "field_qc_status",
            "field_failures",
            "missing_field_metadata",
            "horizontal_residual",
            "vertical_residual",
        ]
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(reshoot_rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Accepted Control"
    accepted_fields = [
        "control_id",
        "northing",
        "easting",
        "elevation",
        "code",
        "horizontal_residual",
        "vertical_residual",
        "source_observations",
        "numbering_flags",
        "field_qc_status",
        "minimum_time_gap_minutes",
        "qc_status",
    ]
    ws.append(accepted_fields)
    for row in accepted_rows:
        ws.append([row.get(k) for k in accepted_fields])
    rs = wb.create_sheet("Reshoot List")
    reshoot_fields = [
        "control_id",
        "reason",
        "existing_shots",
        "next_shots",
        "numbering_flags",
        "field_qc_status",
        "field_failures",
        "missing_field_metadata",
        "horizontal_residual",
        "vertical_residual",
    ]
    rs.append(reshoot_fields)
    for row in reshoot_rows:
        rs.append([row.get(k) for k in reshoot_fields])
    meta = wb.create_sheet("QC Run")
    meta.append(["Run ID", result.get("run_id", "")])
    meta.append(["Horizontal tolerance", result.get("horizontal_tolerance")])
    meta.append(["Vertical tolerance", result.get("vertical_tolerance")])
    meta.append(["Accepted", len(accepted_rows)])
    meta.append(["Reshoot/Review", len(reshoot_rows)])
    context = result.get("coordinate_context") or {}
    meta.append(["Project CRS", context.get("crs", "")])
    meta.append(["Horizontal units", context.get("horizontal_units", "")])
    meta.append(["Vertical units", context.get("vertical_units", "")])
    site = context.get("local_site") or {}
    meta.append(["Local Site enabled", bool(site.get("enabled"))])
    meta.append(["Local Site name", site.get("name", "")])
    meta.append(["Grid-to-ground factor", site.get("grid_to_ground_factor", 1.0)])
    meta.append(["Rotation clockwise deg", site.get("rotation_deg", 0.0)])
    field = result.get("field_requirements") or {}
    meta.append(["Field metadata enforced", bool(field.get("enforced"))])
    meta.append(["Minimum shot separation (minutes)", field.get("min_time_separation_minutes")])
    meta.append(["Minimum epochs", field.get("min_epochs")])
    meta.append(
        [
            "Minimum observation duration (minutes)",
            (field.get("min_duration_seconds") or 0) / 60
            if field.get("min_duration_seconds") is not None
            else "",
        ]
    )
    meta.append(["Minimum satellites", field.get("min_satellites")])
    wb.save(xlsx_path)
    return {
        "accepted_csv": str(accepted_path),
        "reshoot_csv": str(reshoot_path),
        "xlsx": str(xlsx_path),
        "accepted_count": len(accepted_rows),
        "reshoot_count": len(reshoot_rows),
    }


def write_ron_control_deliverables(
    result: dict, source_point_ids: list[str], reports_dir: Path
) -> dict:
    """Create Ronald's plain-text final-control or reshoot deliverables.

    The workbook arithmetic is already represented by ``solve(..., ron_spreadsheet)``.
    This helper only formats those validated results; it does not alter any survey math.
    """
    import re
    from pathlib import Path

    folder = Path(reports_dir)
    folder.mkdir(parents=True, exist_ok=True)
    control_id = str(result.get("control_id") or "Control")
    safe = re.sub("[^A-Za-z0-9._ -]+", "_", control_id).strip(" .") or "Control"
    stamp = utc_now().replace(":", "-")
    qc_path = folder / f"Ron_Control_{safe}_{stamp}_QC.txt"
    residuals = list(result.get("residuals") or [])
    source_ids = [str(x) for x in source_point_ids]
    lines = [
        "SurveySync - Ron 3-Point Control Average / QC",
        f"Control ID: {control_id}",
        "Method: Ron 3-Point Workbook",
        f"Solution ID: {result.get('solution_id', '')}",
        f"Revision: {result.get('revision', '')}",
        f"Timestamp UTC: {utc_now()}",
        "",
        "FINAL AVERAGE",
        f"Northing: {float(result.get('northing')):.4f}",
        f"Easting:  {float(result.get('easting')):.4f}",
        f"Elevation:{float(result.get('elevation')):.4f}"
        if result.get("elevation") is not None
        else "Elevation: N/A",
        f"Horizontal tolerance: {float(result.get('horizontal_tolerance') or 0):.4f}",
        f"Vertical tolerance:   {float(result.get('vertical_tolerance') or 0):.4f}",
        f"Overall status: {('PASS' if result.get('pass') else 'RESHOOT REQUIRED')}",
        "",
        "SOURCE SHOTS",
        "PointID\tHorizontal Residual\tVertical Residual\tStatus",
    ]
    failed: list[tuple[str, dict]] = []
    for idx, residual in enumerate(residuals):
        pid = source_ids[idx] if idx < len(source_ids) else f"Shot {idx + 1}"
        h = float(residual.get("horizontal") or 0)
        dz = residual.get("dz")
        status = "PASS" if residual.get("pass") else "RESHOOT"
        lines.append(f"{pid}\t{h:.4f}\t{('' if dz is None else f'{float(dz):.4f}')}\t{status}")
        if not residual.get("pass"):
            failed.append((pid, residual))
    qc_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    paths: dict[str, str] = {"qc_txt": str(qc_path)}
    if result.get("pass"):
        final_path = folder / f"Ron_Control_{safe}_{stamp}_FINAL.txt"
        final_lines = [
            "SurveySync Final Control",
            f"Control ID: {control_id}",
            f"Northing: {float(result.get('northing')):.4f}",
            f"Easting: {float(result.get('easting')):.4f}",
            f"Elevation: {float(result.get('elevation')):.4f}"
            if result.get("elevation") is not None
            else "Elevation: N/A",
            "Method: Ron 3-Point Workbook",
            "Source PointIDs: " + ", ".join(source_ids),
            f"Solution ID: {result.get('solution_id', '')}",
            f"Revision: {result.get('revision', '')}",
            f"Timestamp UTC: {utc_now()}",
        ]
        final_path.write_text("\n".join(final_lines) + "\n", encoding="utf-8")
        paths["final_control_txt"] = str(final_path)
    else:
        reshoot_path = folder / f"Ron_Control_{safe}_{stamp}_RESHOOT.txt"
        reshoot_lines = [
            "SurveySync GPS Reshoot List",
            f"Control ID: {control_id}",
            "Reason: one or more source shots exceed the selected Ron 3-point QC tolerance.",
            "",
            "PointID\tHorizontal Residual\tVertical Residual\tReason",
        ]
        for pid, residual in failed:
            h = float(residual.get("horizontal") or 0)
            dz = residual.get("dz")
            reasons = []
            if h > float(result.get("horizontal_tolerance") or 0):
                reasons.append("horizontal tolerance exceeded")
            if dz is not None and abs(float(dz)) > float(result.get("vertical_tolerance") or 0):
                reasons.append("vertical tolerance exceeded")
            reshoot_lines.append(
                f"{pid}\t{h:.4f}\t{('' if dz is None else f'{float(dz):.4f}')}\t{'; '.join(reasons) or 'QC failed'}"
            )
        reshoot_path.write_text("\n".join(reshoot_lines) + "\n", encoding="utf-8")
        paths["reshoot_txt"] = str(reshoot_path)
    return paths
