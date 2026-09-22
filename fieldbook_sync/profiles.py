from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Iterable, List, Sequence

from .models import CodeProfile, CodeRule, MatchMode


DEFAULT_PROFILE_NAMES = ["MSD", "MoDOT", "Ameren", "IDOT", "Lambert"]


_HEADER_ALIASES = {
    "code": {"code", "survey code", "feature code", "field code", "point code"},
    "category": {"category", "description", "type", "structure type", "feature type"},
    "include": {"include", "included", "enabled", "use", "active"},
    "match": {"match", "match mode", "matching", "match type"},
}


def _header_key(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").replace("-", " ").split())


def _map_headers(headers: Sequence[object]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for idx, raw in enumerate(headers):
        key = _header_key(raw)
        if not key:
            continue
        for canonical, aliases in _HEADER_ALIASES.items():
            if key in aliases and canonical not in mapped:
                mapped[canonical] = idx
                break
    return mapped


def _parse_include(value: object) -> bool:
    if value is None or str(value).strip() == "":
        return True
    normalized = str(value).strip().lower()
    return normalized not in {"0", "false", "no", "n", "exclude", "excluded", "off"}


def _parse_match(value: object, row_num: int) -> MatchMode:
    if value is None or str(value).strip() == "":
        return MatchMode.EXACT
    normalized = str(value).strip().lower().replace("-", "_").replace(" ", "_")
    aliases = {
        "exact": MatchMode.EXACT,
        "equals": MatchMode.EXACT,
        "equal": MatchMode.EXACT,
        "startswith": MatchMode.STARTS_WITH,
        "starts_with": MatchMode.STARTS_WITH,
        "prefix": MatchMode.STARTS_WITH,
        "contains": MatchMode.CONTAINS,
        "contain": MatchMode.CONTAINS,
    }
    if normalized not in aliases:
        raise ValueError(f"Invalid Match value on row {row_num}: {value}")
    return aliases[normalized]


def _profile_from_rows(name: str, client: str, rows: Iterable[Sequence[object]], *, source_label: str) -> CodeProfile:
    iterator = iter(rows)
    header = next(iterator, None)
    if header is None:
        raise ValueError(f"The {source_label} does not contain a header row.")
    mapped = _map_headers(header)
    if "code" not in mapped:
        raise ValueError(
            f"{source_label.capitalize()} must include a Code column. "
            "Accepted headings include Code, Survey Code, Feature Code, Field Code, or Point Code."
        )

    rules: List[CodeRule] = []
    seen: set[tuple[str, str, bool, MatchMode]] = set()
    for row_num, row in enumerate(iterator, start=2):
        row = list(row)
        def cell(key: str) -> object:
            idx = mapped.get(key)
            return row[idx] if idx is not None and idx < len(row) else None

        raw_code = cell("code")
        if raw_code is None:
            continue
        code = str(raw_code).strip()
        if not code:
            continue
        category = str(cell("category") or "Structure").strip() or "Structure"
        include = _parse_include(cell("include"))
        match = _parse_match(cell("match"), row_num)
        signature = (code.casefold(), category.casefold(), include, match)
        if signature in seen:
            continue
        seen.add(signature)
        rules.append(CodeRule(code=code, category=category, include=include, match=match))

    if not rules:
        raise ValueError(f"The {source_label} contains no usable code rows.")
    return CodeProfile(name=name, client=client or name, codes=rules)


def ensure_default_profiles(profile_dir: Path) -> None:
    profile_dir.mkdir(parents=True, exist_ok=True)
    for name in DEFAULT_PROFILE_NAMES:
        path = profile_dir / f"{safe_filename(name)}.json"
        if not path.exists():
            profile = CodeProfile(
                name=name,
                client=name,
                notes="Add the client's actual storm/sewer survey codes before processing.",
                codes=[],
            )
            save_profile(profile_dir, profile)

    demo_path = profile_dir / "Demo_Profile.json"
    if not demo_path.exists():
        demo = CodeProfile(
            name="Demo Profile",
            client="Demo",
            notes="Example only. Replace these with the client's actual code list.",
            codes=[
                CodeRule(code="DBL GI", category="Double Grate Inlet", include=True, match=MatchMode.EXACT),
                CodeRule(code="GI", category="Grate Inlet", include=True, match=MatchMode.EXACT),
                CodeRule(code="STM MH", category="Storm Manhole", include=True, match=MatchMode.EXACT),
                CodeRule(code="MH", category="Manhole", include=True, match=MatchMode.EXACT),
            ],
        )
        save_profile(profile_dir, demo)


def safe_filename(value: str) -> str:
    cleaned = "".join(c if c.isalnum() or c in "-_ " else "_" for c in value).strip()
    return cleaned.replace(" ", "_") or "profile"


def list_profiles(profile_dir: Path) -> List[CodeProfile]:
    ensure_default_profiles(profile_dir)
    result: List[CodeProfile] = []
    for path in sorted(profile_dir.glob("*.json")):
        try:
            result.append(CodeProfile.model_validate_json(path.read_text(encoding="utf-8")))
        except Exception:
            continue
    return result


def get_profile(profile_dir: Path, name: str) -> CodeProfile:
    profiles = list_profiles(profile_dir)
    for profile in profiles:
        if profile.name == name:
            return profile
    raise KeyError(f"Profile not found: {name}")


def save_profile(profile_dir: Path, profile: CodeProfile, original_name: str | None = None) -> Path:
    """Save a profile, optionally renaming an existing logical profile.

    ``original_name`` identifies the profile currently being edited. When it differs
    from ``profile.name``, the old profile file is removed only after the new file is
    written successfully. The function also refuses to overwrite a different logical
    profile that happens to sanitize to the same filename.
    """
    profile_dir.mkdir(parents=True, exist_ok=True)
    original_name = (original_name or profile.name).strip()
    target = profile_dir / f"{safe_filename(profile.name)}.json"

    existing_profiles: list[tuple[Path, CodeProfile]] = []
    for existing in profile_dir.glob("*.json"):
        try:
            other = CodeProfile.model_validate_json(existing.read_text(encoding="utf-8"))
        except Exception:
            continue
        existing_profiles.append((existing, other))

    # A rename must not silently overwrite another logical profile.
    if original_name != profile.name:
        for existing, other in existing_profiles:
            if other.name == profile.name and other.name != original_name:
                raise ValueError(f"A profile named '{profile.name}' already exists.")

    # Two different logical names can sanitize to the same filename (for example,
    # punctuation variants). Never overwrite one of those by accident.
    if target.exists():
        try:
            target_profile = CodeProfile.model_validate_json(target.read_text(encoding="utf-8"))
        except Exception:
            target_profile = None
        if target_profile is not None and target_profile.name not in {profile.name, original_name}:
            raise ValueError(
                f"Profile name '{profile.name}' conflicts with existing profile "
                f"'{target_profile.name}' after filename sanitization. Choose a different name."
            )

    # Write atomically so a failed write never destroys the original profile.
    temp = target.with_suffix(".tmp")
    temp.write_text(profile.model_dump_json(indent=2), encoding="utf-8")
    temp.replace(target)

    # Remove the old logical profile after the replacement is safely on disk. Also
    # clean up stale duplicate files for the same new logical name.
    for existing, other in existing_profiles:
        should_remove_old = original_name != profile.name and other.name == original_name
        should_remove_stale_new = other.name == profile.name and existing != target
        if (should_remove_old or should_remove_stale_new) and existing != target:
            existing.unlink(missing_ok=True)

    return target


def delete_profile(profile_dir: Path, name: str) -> None:
    for path in profile_dir.glob("*.json"):
        try:
            profile = CodeProfile.model_validate_json(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if profile.name == name:
            path.unlink(missing_ok=True)
            return
    raise KeyError(f"Profile not found: {name}")


def import_profile_csv(name: str, client: str, raw: bytes) -> CodeProfile:
    """Import a profile CSV. Only a code column is required."""
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return _profile_from_rows(name, client, reader, source_label="profile CSV")


def import_profile_excel(name: str, client: str, raw: bytes) -> CodeProfile:
    """Import code rules from the first usable worksheet in an Excel .xlsx-family workbook.

    The importer looks through the first 20 rows of each worksheet for a recognized Code
    heading, so workbooks may contain a title or note rows above the table.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Excel import support is not installed. Run FieldBook Sync setup/repair so openpyxl is installed."
        ) from exc

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read the Excel workbook: {exc}") from exc

    try:
        for ws in wb.worksheets:
            buffered: list[tuple[object, ...]] = []
            row_iter = ws.iter_rows(values_only=True)
            header_pos = None
            for pos, row in enumerate(row_iter, start=1):
                values = tuple(row)
                buffered.append(values)
                if "code" in _map_headers(values):
                    header_pos = pos
                    break
                if pos >= 20:
                    break
            if header_pos is None:
                continue

            header = buffered[-1]
            remaining = (tuple(row) for row in row_iter)
            try:
                return _profile_from_rows(
                    name, client, [header, *remaining], source_label=f"Excel sheet '{ws.title}'"
                )
            except ValueError as exc:
                if "no usable code rows" in str(exc).lower():
                    continue
                raise
    finally:
        wb.close()

    raise ValueError(
        "No worksheet contained a recognized Code column with usable rows. "
        "Accepted headings include Code, Survey Code, Feature Code, Field Code, or Point Code."
    )


def import_profile_file(name: str, client: str, raw: bytes, filename: str) -> CodeProfile:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        return import_profile_csv(name, client, raw)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return import_profile_excel(name, client, raw)
    if suffix == ".xls":
        raise ValueError("Legacy .xls files are not supported. Save the workbook as .xlsx and import it again.")
    raise ValueError("Code-profile import supports .csv, .xlsx, .xlsm, .xltx, and .xltm files.")


def export_profile_csv(profile: CodeProfile) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["Code", "Category", "Include", "Match"])
    for rule in profile.codes:
        writer.writerow([rule.code, rule.category, "Yes" if rule.include else "No", rule.match.value])
    return output.getvalue().encode("utf-8-sig")
