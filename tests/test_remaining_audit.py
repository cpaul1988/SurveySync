"""Regression coverage for G01-G05 desktop-completion safety and build identity."""
import base64
import datetime as dt
import hashlib
import json
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

import pytest
from openpyxl import load_workbook
from test_audit_reproductions import workspace
from test_v940_pointcloud_workflows import _write_minimal_las, _project
from test_v940_crs_report_gis_bridges import _template, _insert_points
from surveysync import workflow_engine as wf
from surveysync import report_template_mapper as tm
from surveysync import updater
from surveysync.release_identity import release_order
from surveysync.manifest_trust import decode_manifest


def test_changed_workflow_cannot_substitute_approved_action(tmp_path, monkeypatch):
    p = _project(tmp_path)
    saved = wf.save_workflow(p, {'name':'Delivery', 'actions':[{'type':'build_deliverable','params':{'label':'original'}}]})
    run = wf.start_workflow(p, saved['workflow_id'])
    called=[]
    monkeypatch.setattr(wf,'_execute_action',lambda *a,**k:called.append(True))
    wf.save_workflow(p, {**saved,'actions':[{'type':'send_notification'}]})
    with pytest.raises(wf.WorkflowError,match='changed'):
        wf.approve_run(p,run['run_id'],approved=True)
    assert not called
    assert wf.approve_run(p,run['run_id'],approved=False)['status']=='REJECTED'


def test_approval_is_not_double_executed(tmp_path, monkeypatch):
    p=_project(tmp_path)
    w=wf.save_workflow(p,{'name':'Delivery','actions':[{'type':'build_deliverable'}]})
    run=wf.start_workflow(p,w['workflow_id']);calls=[]
    monkeypatch.setattr(wf,'_execute_action',lambda *a,**k:calls.append(True) or {})
    def approve(_):
        try:return wf.approve_run(p,run['run_id'],approved=True)['status']
        except wf.WorkflowError:return 'REFUSED'
    with ThreadPoolExecutor(max_workers=2) as pool: results=list(pool.map(approve,range(2)))
    assert sorted(results)==['COMPLETED','REFUSED'] and len(calls)==1


@pytest.mark.parametrize('text',['workflows: [true]', 'workflows: [{id: a, name: A, actions: [{type: run_qa}]}, {id: a, name: B, actions: [{type: run_qa}]}]', 'workflows: [{name: A, enabled: "false", actions: [{type: run_qa}]}]'])
def test_bad_yaml_does_not_replace_definitions(tmp_path,text):
    p=_project(tmp_path);wf.save_workflow(p,{'name':'Keep','actions':[{'type':'run_qa'}]})
    before=wf.export_workflows_yaml(p)
    with pytest.raises((wf.WorkflowError,ValueError)):wf.import_workflows_yaml(p,text)
    assert wf.export_workflows_yaml(p)==before


@pytest.mark.parametrize('endpoint,body',[('/api/v9/pointcloud/import',{'file_path':'missing.las'}),('/api/v9/workflows',{'name':'Test','actions':[{'type':'run_qa'}]}),('/api/v9/reports/templates',{'file_path':'missing.xlsx'}),('/api/v9/crs/project-diagnostics',{})])
def test_stale_panel_refuses_new_project(workspace,endpoint,body):
    client,_,_=workspace
    reply=client.post(endpoint,json=body,headers={'X-SurveySync-Project':'not-current'})
    assert reply.status_code==409,reply.text


def test_truncated_las_rejected(tmp_path):
    from surveysync.pointcloud import inspect_point_cloud, PointCloudError
    file=_write_minimal_las(tmp_path/'bad.las');file.write_bytes(file.read_bytes()[:-1])
    with pytest.raises(PointCloudError,match='truncated'):inspect_point_cloud(file)


def test_template_spaces_formula_protection_and_no_overwrite(tmp_path):
    p=_project(tmp_path);_insert_points(p);path=_template(tmp_path/'original.xlsx')
    w=load_workbook(path);w['Cover']['B1']='{{ project.name }}';w['Cover']['D5']='=1+1';w.save(path);w.close()
    before=path.read_bytes();registered=tm.register_excel_template(p,path)
    with pytest.raises(tm.ReportTemplateError,match='formula'):
        tm.save_template_mapping(p,registered['template_id'],{'scalar_cells':[{'field':'project.name','sheet':'Cover','cell':'D5'}]})
    a=tm.render_excel_template(p,registered['template_id']);b=tm.render_excel_template(p,registered['template_id'])
    assert a['output_path']!=b['output_path'] and path.read_bytes()==before
    w=load_workbook(a['output_path']);assert w['Cover']['B1'].value==p.manifest['name'];assert w['Cover']['D5'].value=='=1+1';w.close()
    with pytest.raises(tm.ReportTemplateError,match='exists'):
        tm.render_excel_template(p,registered['template_id'],output_path=path)
    assert path.read_bytes()==before


def test_template_point_strings_are_not_executable_formulas(tmp_path):
    p=_project(tmp_path);_insert_points(p)
    with p.db.connect() as c:c.execute("UPDATE canonical_points SET description='=1+1'")
    entry=tm.register_excel_template(p,_template(tmp_path/'input.xlsx'),mapping={'point_table':{'sheet':'Points','start_row':2,'columns':{'A':'point_id','E':'description'}}})
    out=tm.render_excel_template(p,entry['template_id']);w=load_workbook(out['output_path'])
    assert w['Points']['E2'].value=='=1+1' and w['Points']['E2'].data_type=='s';w.close()


@pytest.mark.parametrize('kwargs',[{'sample_x':1}, {'sample_x':float('inf'),'sample_y':0}, {'sample_x':0,'sample_y':1000}, {'area_of_interest':{'west':10,'east':-10,'south':0,'north':1}}])
def test_invalid_crs_samples_return_errors(kwargs):
    from surveysync.crs_diagnostics import operation_diagnostics,CrsDiagnosticError
    with pytest.raises(CrsDiagnosticError):operation_diagnostics('EPSG:4326','EPSG:3857',**kwargs)


def test_grass_module_requires_literal_separator(tmp_path):
    from surveysync.gis_bridges import run_grass_module,GisBridgeError
    with pytest.raises(GisBridgeError,match='standard'):run_grass_module(_project(tmp_path),'gXremove',{})


@pytest.mark.parametrize('left,right',[('9.4.2-beta.2','9.4.2-beta.10'),('9.4.2-alpha','9.4.2-beta'),('9.4.2-beta.10','9.4.2-rc.1'),('9.4.2-rc.1','9.4.2'),('9.4.2','9.4.3-beta.1')])
def test_release_order(left,right):assert release_order(left)<release_order(right)


@pytest.mark.parametrize('bad',['9.4.2-beta.01','9.4.2-beta..1','9.4.2/evil','9.4.2-','9.4.2+','999999.1'])
def test_malformed_identity_rejected(bad):
    with pytest.raises(ValueError):release_order(bad)


def test_build_metadata_not_precedence():assert release_order('9.4.2-beta.2+first')==release_order('9.4.2-beta.2+second')


def test_beta_check_and_same_artifact_promotion(tmp_path,monkeypatch):
    monkeypatch.setattr(updater, "__version__", "9.4.1")
    from surveysync.config import ConfigStore
    store=ConfigStore(tmp_path/'config');cfg=store.load();cfg.release_channel='beta';store.save(cfg)
    monkeypatch.setattr(updater,'installed_release_id',lambda _: '9.4.1-beta.2')
    rel={'version':'9.4.1','release_id':'9.4.1-beta.10','installer_url':'https://example.com/test.exe','sha256':'a'*64,'size_bytes':100}
    monkeypatch.setattr(updater,'fetch_manifest',lambda _: {'channels':{'beta':rel,'stable':{**rel,'release_id':'9.4.1'}}})
    assert updater.check(store)['update_available']
    monkeypatch.setattr(updater,'installed_release_id',lambda _: '9.4.1-beta.10')
    assert not updater.check(store)['update_available']
    (store.root/'installed_update.json').write_text(json.dumps({'version':'9.4.1','release_id':'9.4.1-beta.10','sha256':'a'*64}))
    cfg.release_channel='stable';store.save(cfg)
    assert not updater.check(store)['update_available']


def signed_fixture():
    pytest.importorskip('cryptography')
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    from cryptography.hazmat.primitives.serialization import Encoding,PublicFormat
    key=Ed25519PrivateKey.generate() # Disposable test key, never a production signing identity.
    payload=json.dumps({'product':'SurveySync','expires_utc':'2099-01-01T00:00:00Z','channels':{}}).encode()
    enc=lambda b:base64.b64encode(b).decode()
    envelope={'schema':'surveysync.signed-manifest.v1','key_id':'fixture','payload':enc(payload),'signature':enc(key.sign(payload))}
    policy={'require_signature':True,'keys':{'fixture':enc(key.public_key().public_bytes(Encoding.Raw,PublicFormat.Raw))}}
    return envelope,policy


def test_signed_manifest_and_expiry():
    e,p=signed_fixture();assert decode_manifest(json.dumps(e).encode(),p)['product']=='SurveySync'
    with pytest.raises(ValueError,match='expired'):decode_manifest(json.dumps(e).encode(),p,now=dt.datetime(2100,1,1,tzinfo=dt.timezone.utc))


@pytest.mark.parametrize('attack',['payload','signature','key_id','unsigned'])
def test_signature_tampering_fails_closed(attack):
    e,p=signed_fixture()
    if attack=='unsigned':e={'product':'SurveySync'}
    elif attack=='key_id':e[attack]='attacker'
    else:e[attack]=base64.b64encode(b'altered').decode()
    with pytest.raises(ValueError):decode_manifest(json.dumps(e).encode(),p)


def test_promoted_physical_beta_stamp_does_not_reinstall(tmp_path,monkeypatch):
    monkeypatch.setattr(updater, "__version__", "9.4.1")
    from surveysync.config import ConfigStore
    store=ConfigStore(tmp_path/'config');cfg=store.load();cfg.release_channel='stable';store.save(cfg)
    monkeypatch.setattr(updater,'installed_release_id',lambda _: '9.4.1-beta.10')
    rel={'version':'9.4.1','release_id':'9.4.1-beta.10','promoted_from':'beta','artifact_identity':'sha256:'+'a'*64,
         'installer_url':'https://example.com/setup.exe','sha256':'a'*64,'size_bytes':100}
    monkeypatch.setattr(updater,'fetch_manifest',lambda _: {'channels':{'stable':rel}})
    assert not updater.check(store)['update_available']
    rel.pop('promoted_from')
    with pytest.raises(ValueError,match='explicit'):updater.check(store)


def test_recursive_and_nonfinite_yaml_params_rejected_before_save(tmp_path):
    p=_project(tmp_path)
    for text in ('workflows: [{name: bad, actions: [{type: run_qa, params: {bad: .nan}}]}]',
                 'workflows: [{name: bad, actions: [{type: run_qa, params: &a {self: *a}}]}]'):
        with pytest.raises(wf.WorkflowError):wf.import_workflows_yaml(p,text)
    assert wf.list_workflows(p)==[]
