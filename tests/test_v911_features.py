from __future__ import annotations

from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fieldbook_sync.models import PipeMeasurement, ResultRecord
from fieldbook_sync.intelligence import infer_network, apply_survey_qc
from surveysync.control import write_ron_control_deliverables
from surveysync.reports import crew_range_recommendations, crew_ranges_txt, write_crew_ranges_xlsx
from surveysync.trimble_job import parse_jobxml_points, numeric_point_ids


def _sample_jxl(path: Path) -> Path:
    path.write_text(
        '''<?xml version="1.0" encoding="utf-8"?>
<JobXML jobName="BRT Test" version="6.6" product="Trimble Access" productVersion="2026.10">
  <Environment><CoordinateSystem><CoordinateSystemName>Texas South Central</CoordinateSystemName></CoordinateSystem></Environment>
  <FieldBook><PointRecord/><PointRecord/></FieldBook>
  <Reductions>
    <Point><Name>2500</Name><Code>SSMH</Code><SurveyMethod>GNSS</SurveyMethod><Grid><North>1000.0</North><East>2000.0</East><Elevation>10.0</Elevation></Grid></Point>
    <Point><Name>2801</Name><Code>SSMH</Code><Grid><North>1100.0</North><East>2100.0</East><Elevation>11.0</Elevation></Grid></Point>
    <Point><Name>CTRL-A</Name><Code>CP</Code><Grid><North>1200.0</North><East>2200.0</East><Elevation>12.0</Elevation></Grid></Point>
  </Reductions>
</JobXML>''',
        encoding="utf-8",
    )
    return path


def test_trimble_jobxml_points_and_numeric_ranges(tmp_path):
    parsed = parse_jobxml_points(_sample_jxl(tmp_path / "brt.jxl"))
    assert parsed["metadata"]["job_name"] == "BRT Test"
    assert parsed["metadata"]["point_count"] == 3
    assert parsed["points"][0]["point_id"] == "2500"
    assert parsed["points"][0]["northing"] == 1000.0
    ids, ignored = numeric_point_ids(parsed["points"])
    assert ids == [2500, 2801]
    assert ignored == 1


def test_point_range_txt_and_xlsx_outputs(tmp_path):
    summary = crew_range_recommendations([1, 100, 2500, 2801, 4000], min_capacity=100, sort_order="ascending")
    text = crew_ranges_txt(summary)
    assert "From" in text and "To" in text
    assert "and above" in text
    output = write_crew_ranges_xlsx(summary, tmp_path / "ranges.xlsx")
    assert output.is_file()
    ws = load_workbook(output, read_only=True)["Available Point Ranges"]
    assert ws["A7"].value == "From"
    assert ws["B7"].value == "To"


def test_ron_control_deliverables_pass_and_reshoot(tmp_path):
    base = {
        "control_id": "CP-1", "solution_id": "s1", "revision": 1,
        "northing": 1000.0, "easting": 2000.0, "elevation": 100.0,
        "horizontal_tolerance": 0.10, "vertical_tolerance": 0.10,
        "residuals": [
            {"horizontal": 0.02, "dz": 0.01, "pass": True},
            {"horizontal": 0.03, "dz": -0.02, "pass": True},
            {"horizontal": 0.01, "dz": 0.01, "pass": True},
        ],
        "pass": True,
    }
    passed = write_ron_control_deliverables(base, ["100", "101", "102"], tmp_path)
    assert Path(passed["qc_txt"]).is_file()
    assert Path(passed["final_control_txt"]).is_file()
    assert "reshoot_txt" not in passed

    failed = dict(base)
    failed["pass"] = False
    failed["residuals"] = [
        {"horizontal": 0.02, "dz": 0.01, "pass": True},
        {"horizontal": 0.15, "dz": 0.02, "pass": False},
        {"horizontal": 0.01, "dz": 0.25, "pass": False},
    ]
    reshoot = write_ron_control_deliverables(failed, ["100", "101", "102"], tmp_path)
    txt = Path(reshoot["reshoot_txt"]).read_text(encoding="utf-8")
    assert "101" in txt and "horizontal tolerance exceeded" in txt
    assert "102" in txt and "vertical tolerance exceeded" in txt


def test_brt_explicit_to_point_topology_beats_geometric_guess():
    src = ResultRecord(
        point_id="3389", northing=0.0, easting=0.0, elevation=10.0,
        code="SSMH", status="YES", dip_status="YES",
        pipes=[PipeMeasurement(azimuth_deg=90.0, leader_direction_deg=90.0, connected_point_raw="Az = to 3394")],
    )
    target = ResultRecord(
        point_id="3394", northing=0.0, easting=100.0, elevation=9.0,
        code="SSMH", status="YES", dip_status="YES",
        pipes=[PipeMeasurement(azimuth_deg=270.0)],
    )
    closer_wrong = ResultRecord(
        point_id="3390", northing=0.0, easting=20.0, elevation=9.5,
        code="SSMH", status="YES", dip_status="YES", pipes=[]
    )
    edges = infer_network([src, closer_wrong, target], max_distance=1500, max_bearing_error=25)
    explicit = next(e for e in edges if e.from_point == "3389")
    assert explicit.to_point == "3394"
    assert explicit.status == "STRONG"
    assert src.pipes[0].connected_point_id == "3394"


def test_brt_written_azimuth_vs_leader_direction_routes_review():
    r = ResultRecord(
        point_id="1", northing=0, easting=0, elevation=10, code="MH",
        pipes=[PipeMeasurement(azimuth_deg=10.0, leader_direction_deg=100.0)]
    )
    apply_survey_qc([r])
    assert any("leader direction disagree" in flag for flag in r.qc_flags)


def test_reportsync_point_range_defaults_to_current_project_and_accepts_jxl(tmp_path, monkeypatch):
    monkeypatch.setenv("SURVEYSYNC_CONFIG_ROOT", str(tmp_path / "cfg"))
    from fieldbook_sync.app import app
    from surveysync import router as r
    from surveysync.audit import utc_now

    r.config_store = r.ConfigStore(tmp_path / "cfg")
    r.current_project = None
    client = TestClient(app)
    created = client.post('/api/v9/project/create', json={
        "parent_folder": str(tmp_path), "name": "Range Project", "crs": "EPSG:2278",
        "horizontal_units": "us_survey_feet", "vertical_units": "us_survey_feet"
    })
    assert created.status_code == 200, created.text
    p = r.current_project
    assert p is not None
    now = utc_now()
    with p.db.connect() as conn:
        for pid in ("2500", "2801", "4000"):
            conn.execute(
                "INSERT INTO canonical_points(point_uuid,point_id,northing,easting,elevation,description,point_class,source_id,derived_from_json,crs,horizontal_units,vertical_units,review_state,revision,created_utc,modified_utc) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (uuid4().hex, pid, 1000.0, 2000.0, 10.0, "", "survey", None, "[]", "EPSG:2278", "us_survey_feet", "us_survey_feet", "UNREVIEWED", 1, now, now),
            )
    project_result = client.post('/api/v9/reports/point-ranges', json={"source_mode": "project", "min_capacity": 100})
    assert project_result.status_code == 200, project_result.text
    body = project_result.json()
    assert body["used_count"] == 3
    assert body["source_mode"] == "project"
    assert all(Path(x).is_file() for x in body["report_paths"].values())

    jxl = _sample_jxl(tmp_path / "external.jxl")
    file_result = client.post('/api/v9/reports/point-ranges', json={"source_mode": "file", "file_path": str(jxl), "min_capacity": 100})
    assert file_result.status_code == 200, file_result.text
    assert file_result.json()["used_count"] == 2
    assert file_result.json()["ignored_point_ids"] == 1


def test_trimble_job_converter_uses_official_cli_contract(tmp_path, monkeypatch):
    import subprocess
    from surveysync import trimble_job as tj

    source = tmp_path / "crew.job"
    source.write_bytes(b"opaque vendor bytes")
    # This is intentionally only a file-system fixture.  Mock subprocess execution
    # so the test remains portable on Windows, where an extensionless POSIX shebang
    # script is not a valid Win32 executable (WinError 193).
    fake = tmp_path / "AsciiFileGenerator.exe"
    fake.write_bytes(b"test fixture; not executed")
    out = tmp_path / "crew.jxl"
    captured = {}

    def fake_run(command, **kwargs):
        captured["command"] = list(command)
        Path(command[3]).write_text("<JobXML><Reductions/></JobXML>", encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr(tj.subprocess, "run", fake_run)
    result = tj.convert_job_to_jxl(source, out, generator_path=fake)
    assert captured["command"] == [str(fake.resolve()), str(source.resolve()), "Trimble JobXML", str(out.resolve())]
    assert result["return_code"] == 0
    assert out.read_text(encoding="utf-8").startswith("<JobXML>")


def test_fieldbooksync_parses_trimble_jxl_through_code_profile(tmp_path):
    from fieldbook_sync.models import CodeProfile, CodeRule, MatchMode
    from fieldbook_sync.survey import parse_survey_file, extract_numeric_point_ids_file
    profile = CodeProfile(name="BRT", codes=[CodeRule(code="SSMH", category="Sewer", include=True, match=MatchMode.EXACT)])
    jxl = _sample_jxl(tmp_path / "field.jxl")
    points, issues = parse_survey_file(jxl, "field.jxl", profile)
    assert [p.point_id for p in points] == ["2500", "2801"]
    assert all(p.category == "Sewer" for p in points)
    ids, ignored = extract_numeric_point_ids_file(jxl)
    assert ids == [2500, 2801]
    assert ignored == 1


def test_main_trimble_jobxml_import_preserves_source_and_holds_conflicts(tmp_path, monkeypatch):
    monkeypatch.setenv("SURVEYSYNC_CONFIG_ROOT", str(tmp_path / "cfg_trimble"))
    from fieldbook_sync.app import app
    from surveysync import router as r
    from surveysync.audit import utc_now
    r.config_store = r.ConfigStore(tmp_path / "cfg_trimble")
    r.current_project = None
    client = TestClient(app)
    made = client.post('/api/v9/project/create', json={
        "parent_folder": str(tmp_path), "name": "Trimble Project", "crs": "EPSG:2278",
        "horizontal_units": "us_survey_feet", "vertical_units": "us_survey_feet"
    })
    assert made.status_code == 200, made.text
    p = r.current_project
    now = utc_now()
    with p.db.connect() as conn:
        conn.execute(
            "INSERT INTO canonical_points(point_uuid,point_id,northing,easting,elevation,description,point_class,source_id,derived_from_json,crs,horizontal_units,vertical_units,review_state,revision,created_utc,modified_utc) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (uuid4().hex, "2500", 1.0, 2.0, 3.0, "existing", "survey", None, "[]", "EPSG:2278", "us_survey_feet", "us_survey_feet", "UNREVIEWED", 1, now, now),
        )
    jxl = _sample_jxl(tmp_path / "direct.jxl")
    res = client.post('/api/v9/trimble/job-import', json={"file_path": str(jxl), "import_points": True, "point_class": "survey"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["canonical_points_imported"] == 2  # 2801 + CTRL-A; 2500 is held
    assert body["point_id_conflicts"] == ["2500"]
    assert Path(body["stored_source"]).is_file()
    assert Path(body["jobxml_path"]).is_file()
    with p.db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM canonical_points WHERE point_id='2500'").fetchone()[0] == 1


def test_main_direct_job_import_uses_configured_trimble_converter(tmp_path, monkeypatch):
    import subprocess
    from surveysync import trimble_job as tj

    monkeypatch.setenv("SURVEYSYNC_CONFIG_ROOT", str(tmp_path / "cfg_job"))
    fake = tmp_path / "AsciiFileGenerator.exe"
    fake.write_bytes(b"test fixture; not executed")
    monkeypatch.setenv("SURVEYSYNC_TRIMBLE_ASCII_GENERATOR", str(fake))

    def fake_run(command, **kwargs):
        Path(command[3]).write_text(
            '<JobXML jobName="Direct JOB"><Reductions><Point><Name>7001</Name><Code>SSMH</Code>'
            '<Grid><North>10</North><East>20</East><Elevation>5</Elevation></Grid>'
            '</Point></Reductions></JobXML>',
            encoding="utf-8",
        )
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr(tj.subprocess, "run", fake_run)
    source = tmp_path / "crew.job"
    source.write_bytes(b"opaque trimble job fixture")

    from fieldbook_sync.app import app
    from surveysync import router as r
    r.config_store = r.ConfigStore(tmp_path / "cfg_job")
    r.current_project = None
    client = TestClient(app)
    made = client.post('/api/v9/project/create', json={
        "parent_folder": str(tmp_path), "name": "Direct Job Project", "crs": "EPSG:2278",
        "horizontal_units": "us_survey_feet", "vertical_units": "us_survey_feet"
    })
    assert made.status_code == 200, made.text
    res = client.post('/api/v9/trimble/job-import', json={"file_path": str(source), "import_points": True})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["conversion"]["converted"] is True
    assert body["canonical_points_imported"] == 1
    assert Path(body["stored_source"]).suffix.lower() == ".job"
    assert Path(body["jobxml_path"]).suffix.lower() == ".jxl"


def test_v911_shell_exposes_point_file_browse_and_trimble_job_controls(tmp_path, monkeypatch):
    monkeypatch.setenv("SURVEYSYNC_CONFIG_ROOT", str(tmp_path / "cfg_ui_v911"))
    from fieldbook_sync.app import app
    from surveysync import router as r
    r.config_store = r.ConfigStore(tmp_path / "cfg_ui_v911")
    client = TestClient(app)
    html = client.get('/').text
    assert 'id="rangeSource"' in html
    assert 'Current Project Points' in html
    assert 'id="rangeBrowse"' in html
    assert 'id="trimbleJobPath"' in html
    assert 'id="trimbleJobBrowse"' in html
    assert 'id="trimbleJobImport"' in html
    fieldbook = client.get('/fieldbook').text
    assert 'accept=".txt,.csv,.tsv,.pnezd,.asc,.job,.jxl,.xml"' in fieldbook


def test_trimble_jobxml_inventorydata_fallback_when_reductions_empty(tmp_path):
    path = tmp_path / "tbc_inventory.jxl"
    path.write_text(
        """<?xml version="1.0" encoding="utf-8"?>
<JOBFile jobName="TBC Inventory" version="6.33" product="Trimble Business Center" productVersion="43.0">
  <FieldBook />
  <Reductions />
  <InventoryData>
    <Point><ID>1</ID><Name>100A</Name><Code>CTRL</Code><Grid><North>304725.100</North><East>273455.200</East><Elevation>152.700</Elevation></Grid><SurveyMethod>KeyedIn</SurveyMethod></Point>
    <Point><ID>2</ID><Name>100B</Name><Code>CTRL</Code><Grid><North>304725.110</North><East>273455.210</East><Elevation>152.710</Elevation></Grid><SurveyMethod>KeyedIn</SurveyMethod></Point>
    <Point><ID>3</ID><Name>100C</Name><Code>CTRL</Code><Grid><North>304725.090</North><East>273455.195</East><Elevation>152.695</Elevation></Grid><SurveyMethod>KeyedIn</SurveyMethod></Point>
  </InventoryData>
</JOBFile>""",
        encoding="utf-8",
    )
    parsed = parse_jobxml_points(path)
    assert parsed["metadata"]["point_source"] == "InventoryData"
    assert parsed["metadata"]["reduction_point_count"] == 0
    assert parsed["metadata"]["inventory_point_count"] == 3
    assert [p["point_id"] for p in parsed["points"]] == ["100A", "100B", "100C"]
