from __future__ import annotations

import json
import subprocess
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from surveysync.crs_diagnostics import (
    operation_diagnostics,
    project_crs_diagnostics,
)
from surveysync.gis_bridges import (
    GisBridgeError,
    bridge_status,
    qgis_algorithms,
    run_grass_module,
    run_qgis_algorithm,
)
from surveysync.project import SurveyProject
from surveysync.report_template_mapper import (
    inspect_excel_template,
    register_excel_template,
    render_excel_template,
    save_template_mapping,
)


def _project(tmp_path: Path, *, crs: str = "EPSG:2278", units: str = "us_survey_feet"):
    return SurveyProject.create(
        tmp_path,
        "Diagnostics Project",
        crs=crs,
        horizontal_units=units,
        vertical_units=units,
        client="Example Client",
        project_number="26-001",
    )


def _template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    ws["A1"] = "Project"
    ws["B1"] = "{{project.name}}"
    ws["A2"] = "Client"
    ws["B2"] = "Client: {{project.client}}"
    ws["A3"] = "Point Count"
    ws["B3"] = "{{counts.points}}"
    points = wb.create_sheet("Points")
    points["A1"] = "Point"
    points["B1"] = "Northing"
    points["C1"] = "Easting"
    points["D1"] = "Elevation"
    points["A2"] = "template"
    points["B2"] = 0
    points["C2"] = 0
    points["D2"] = 0
    points["B2"].number_format = "0.0000"
    points["C2"].number_format = "0.0000"
    points["D2"].number_format = "0.0000"
    wb.save(path)
    wb.close()
    return path


def _insert_points(project: SurveyProject) -> None:
    from surveysync.audit import utc_now

    now = utc_now()
    with project.db.connect() as conn:
        for pid, north, east, elev in (
            ("100", 1000.0, 2000.0, 10.0),
            ("101", 1001.25, 2001.5, 10.5),
        ):
            conn.execute(
                """INSERT INTO canonical_points(
                    point_uuid,point_id,northing,easting,elevation,description,
                    point_class,source_id,derived_from_json,crs,horizontal_units,
                    vertical_units,review_state,revision,created_utc,modified_utc
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    f"uuid-{pid}",
                    pid,
                    north,
                    east,
                    elev,
                    "CTRL",
                    "survey",
                    None,
                    "[]",
                    project.manifest["crs"],
                    project.manifest["horizontal_units"],
                    project.manifest["vertical_units"],
                    "REVIEWED",
                    1,
                    now,
                    now,
                ),
            )


def test_crs_operation_diagnostics_reports_accuracy_and_area_of_use():
    result = operation_diagnostics(
        "EPSG:4326",
        "EPSG:3857",
        sample_x=-95.0,
        sample_y=29.0,
    )

    assert result["source"]["authority"] == "EPSG:4326"
    assert result["target"]["authority"] == "EPSG:3857"
    assert result["available_operation_count"] >= 1
    assert result["available_operations"][0]["description"]
    assert result["sample"]["inside_source_area_of_use"] is True
    assert result["sample"]["inside_target_area_of_use"] is True
    assert result["sample"]["target"]["x"] != pytest.approx(-95.0)


def test_project_crs_diagnostics_flags_configured_unit_mismatch(tmp_path):
    project = _project(tmp_path, crs="EPSG:3857", units="us_survey_feet")

    result = project_crs_diagnostics(project)

    assert result["status"] == "REVIEW"
    assert result["unit_check"]["matches"] is False
    assert "metre" in result["unit_check"]["crs_unit"].casefold()
    assert "PROJECT_UNIT_MISMATCH" in {
        warning["code"] for warning in result["warnings"]
    }


def test_template_inspection_finds_placeholders(tmp_path):
    path = _template(tmp_path / "client_template.xlsx")

    result = inspect_excel_template(path)

    assert {sheet["name"] for sheet in result["sheets"]} == {"Cover", "Points"}
    assert any(item["cell"] == "B1" for item in result["placeholders"])
    assert any("project.client" in item["fields"] for item in result["placeholders"])


def test_template_mapper_preserves_source_and_renders_draft_workbook(tmp_path):
    project = _project(tmp_path)
    _insert_points(project)
    path = _template(tmp_path / "client_template.xlsx")

    registered = register_excel_template(project, path, name="Client Control Sheet")
    template_id = registered["template_id"]
    save_template_mapping(
        project,
        template_id,
        {
            "scalar_cells": [
                {"field": "project.number", "sheet": "Cover", "cell": "D1"},
            ],
            "replace_placeholders": True,
            "point_table": {
                "sheet": "Points",
                "start_row": 2,
                "columns": {
                    "A": "point_id",
                    "B": "northing",
                    "C": "easting",
                    "D": "elevation",
                },
            },
        },
    )

    rendered = render_excel_template(project, template_id)
    output = Path(rendered["output_path"])
    assert output.is_file()
    assert rendered["deliverable"]["status"] == "DRAFT"
    assert rendered["point_count"] == 2

    wb = load_workbook(output, data_only=False)
    try:
        assert wb["Cover"]["B1"].value == "Diagnostics Project"
        assert wb["Cover"]["B2"].value == "Client: Example Client"
        assert wb["Cover"]["D1"].value == "26-001"
        assert wb["Points"]["A2"].value == "100"
        assert wb["Points"]["A3"].value == "101"
        assert wb["Points"]["B3"].value == pytest.approx(1001.25)
        assert wb["Points"]["B3"].number_format == "0.0000"
    finally:
        wb.close()

    sources = [row for row in project.sources() if row["module"] == "ReportSync"]
    assert sources and sources[0]["original_name"] == "client_template.xlsx"
    stored = project.paths.root / sources[0]["stored_path"]
    original = load_workbook(stored, data_only=False)
    try:
        assert original["Cover"]["B1"].value == "{{project.name}}"
        assert original["Points"]["A2"].value == "template"
    finally:
        original.close()


def test_qgis_bridge_uses_argument_array_and_parses_algorithm_list(tmp_path, monkeypatch):
    fake = tmp_path / "qgis_process.exe"
    fake.write_bytes(b"fixture")

    import surveysync.gis_bridges as bridges

    monkeypatch.setattr(bridges, "find_qgis_process", lambda explicit=None: fake)

    def fake_run(command, **kwargs):
        assert kwargs["shell"] is False
        if command[1] == "list":
            stdout = "native:buffer Buffer\nnative:centroids Centroids\n"
        else:
            stdout = '{"OUTPUT":"out.gpkg"}'
        return subprocess.CompletedProcess(command, 0, stdout=stdout, stderr="")

    monkeypatch.setattr(bridges.subprocess, "run", fake_run)

    listed = qgis_algorithms()
    assert [item["id"] for item in listed["algorithms"]] == [
        "native:buffer",
        "native:centroids",
    ]

    project = _project(tmp_path / "project")
    result = run_qgis_algorithm(
        project,
        "native:buffer",
        {"DISTANCE": 10, "SEGMENTS": 5},
    )
    assert result["return_code"] == 0
    assert result["parsed_output"]["OUTPUT"] == "out.gpkg"
    assert result["command"][1:4] == ["run", "native:buffer", "--"]


def test_qgis_bridge_rejects_invalid_algorithm_id(tmp_path):
    project = _project(tmp_path)

    with pytest.raises(GisBridgeError, match="Invalid QGIS algorithm ID"):
        run_qgis_algorithm(project, "native:buffer & calc.exe", {})


def test_grass_bridge_uses_temp_project_and_validated_module(tmp_path, monkeypatch):
    fake = tmp_path / "grass.exe"
    fake.write_bytes(b"fixture")

    import surveysync.gis_bridges as bridges

    monkeypatch.setattr(bridges, "find_grass", lambda explicit=None: fake)

    captured = {}

    def fake_run(command, **kwargs):
        captured["command"] = list(command)
        assert kwargs["shell"] is False
        return subprocess.CompletedProcess(command, 0, stdout="ok", stderr="")

    monkeypatch.setattr(bridges.subprocess, "run", fake_run)
    project = _project(tmp_path / "project")

    result = run_grass_module(
        project,
        "v.buffer",
        {"distance": 25, "input": "roads"},
        flags=["t"],
    )

    assert result["return_code"] == 0
    assert captured["command"][1:5] == [
        "--tmp-project",
        "EPSG:2278",
        "--exec",
        "v.buffer",
    ]
    assert "-t" in captured["command"]
    assert "distance=25" in captured["command"]


def test_bridge_status_is_optional_when_runtimes_are_missing(monkeypatch):
    import surveysync.gis_bridges as bridges

    monkeypatch.setattr(bridges, "find_qgis_process", lambda explicit=None: None)
    monkeypatch.setattr(bridges, "find_grass", lambda explicit=None: None)

    status = bridge_status()

    assert status["qgis"]["ready"] is False
    assert status["grass"]["ready"] is False
    assert status["embedded_gpl_code"] is False
    assert status["shell_execution"] is False


def test_new_integration_routes_are_registered(tmp_path, monkeypatch):
    monkeypatch.setenv("SURVEYSYNC_CONFIG_ROOT", str(tmp_path / "cfg"))

    from fieldbook_sync import app as field_app
    from surveysync import router as survey_router

    survey_router.config_store = survey_router.ConfigStore(tmp_path / "cfg")
    survey_router.current_project = None
    field_app.runtime = field_app.Runtime(tmp_path / "field_runtime")
    client = TestClient(field_app.app)

    created = client.post(
        "/api/v9/project/create",
        json={
            "parent_folder": str(tmp_path / "projects"),
            "name": "Route Project",
            "crs": "EPSG:2278",
            "horizontal_units": "us_survey_feet",
            "vertical_units": "us_survey_feet",
        },
    )
    assert created.status_code == 200, created.text

    crs = client.post(
        "/api/v9/crs/project-diagnostics",
        json={"target_crs": "EPSG:4326"},
    )
    assert crs.status_code == 200, crs.text

    templates = client.get("/api/v9/reports/templates")
    assert templates.status_code == 200
    assert templates.json()["templates"] == []

    bridges = client.get("/api/v9/gis-bridges/status")
    assert bridges.status_code == 200
    assert "qgis" in bridges.json() and "grass" in bridges.json()


def test_gis_bridge_audit_events_do_not_store_shell_commands(tmp_path, monkeypatch):
    fake = tmp_path / "qgis_process.exe"
    fake.write_bytes(b"fixture")

    import surveysync.gis_bridges as bridges

    monkeypatch.setattr(bridges, "find_qgis_process", lambda explicit=None: fake)
    monkeypatch.setattr(
        bridges.subprocess,
        "run",
        lambda command, **kwargs: subprocess.CompletedProcess(
            command, 0, stdout="done", stderr=""
        ),
    )

    project = _project(tmp_path / "project")
    run_qgis_algorithm(project, "native:centroids", {"INPUT": "roads.gpkg"})

    events = [
        row
        for row in project.db.recent_audit(20)
        if row["action"] == "QGIS_PROCESS_RUN"
    ]
    assert events
    details = events[0]["details"]
    assert details["algorithm_id"] == "native:centroids"
    assert "shell" not in json.dumps(details).lower()
