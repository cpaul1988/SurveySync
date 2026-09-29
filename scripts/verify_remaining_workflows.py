"""Actual local web application controls for G01-G05; synthetic data only.

No external emails, provider requests, real projects or publication. Optional
laspy/lazrs is exercised when installed; unavailable capability stays explicit.
"""
from __future__ import annotations
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'remaining-evidence'
BASE = 'http://127.0.0.1:18771'
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'tests'))
from test_v940_pointcloud_workflows import _write_minimal_las
from test_v940_crs_report_gis_bridges import _template


def api(path, data=None):
    req = urllib.request.Request(BASE+path, data=None if data is None else json.dumps(data).encode(), headers={'Content-Type':'application/json'})
    with urllib.request.urlopen(req, timeout=30) as response:
        return json.load(response)


def main():
    OUT.mkdir(exist_ok=True)
    report = {'status':'FAIL','checks':[], 'optional':{}}
    with tempfile.TemporaryDirectory(prefix='ss-remaining-') as tmp, (OUT/'server.log').open('w', encoding='utf-8') as log:
        temp = Path(tmp)
        las = _write_minimal_las(temp/'sample.las')
        workbook = _template(temp/'template.xlsx')
        w = load_workbook(workbook); w['Cover']['D5']='=1+1'; w.save(workbook); w.close()
        source = temp/'points.csv'
        source.write_text('PointID,Northing,Easting,Elevation,Code\n001A,1000,2000,10,CP\n001B,1001,2001,11,CP\n',encoding='utf-8')
        hashes = {p:hashlib.sha256(p.read_bytes()).hexdigest() for p in (las,workbook,source)}
        server = subprocess.Popen([sys.executable,'-m','uvicorn','fieldbook_sync.app:app','--host','127.0.0.1','--port','18771'], cwd=ROOT, env=dict(os.environ,SURVEYSYNC_CONFIG_ROOT=str(temp/'config'),SURVEYSYNC_FIELD_ROOT=str(temp/'field')), stdout=log,stderr=subprocess.STDOUT)
        try:
            for _ in range(100):
                try:
                    api('/api/v9/release-notes');break
                except (OSError,ValueError):time.sleep(.2)
            else:raise TimeoutError('Synthetic test server unavailable')
            api('/api/v9/project/create',{'parent_folder':str(temp/'projects'),'name':'Remaining Workflow QA','crs':'EPSG:2278','horizontal_units':'us_survey_feet','vertical_units':'us_survey_feet','client':'Synthetic client'})
            api('/api/v9/points/import',{'file_path':str(source)})
            project_id = api('/api/v9/status')['project']['project_id']
            api('/api/v9/config/ui',{'appearance':'dark','theme':'carbon','accent':'default'})
            with sync_playwright() as pw:
                browser=pw.chromium.launch(executable_path=os.environ.get('SURVEYSYNC_TEST_CHROMIUM') or None, args=['--no-sandbox'] if sys.platform!='win32' else [])
                context=browser.new_context(viewport={'width':1488,'height':940})
                context.tracing.start(screenshots=True,snapshots=True)
                page=context.new_page();errors=[]
                page.on('pageerror',lambda e:errors.append(str(e)))
                page.on('dialog',lambda d:d.accept())
                page.route('**/api/v9/update/check',lambda r:r.fulfill(json={'update_available':False,'current_version':'9.4.5','channel':'stable'}))
                def nav(module,label,view):
                    page.locator(f'.module-tab[data-module="{module}"]').click()
                    page.locator('#moduleNav .module-nav-btn').filter(has_text=label).last.click()
                    expect(page.locator('body')).to_have_attribute('data-active-view',view)
                    expect(page.locator('#'+view+' fieldset')).to_be_enabled(timeout=15000)
                    expect(page.locator('#'+view+' .dw-error')).to_have_class('dw-error notice hidden')
                def call(button,endpoint):
                    expect(page.locator('#'+button)).to_be_enabled(timeout=15000)
                    with page.expect_response(lambda r:r.url.split('?')[0].endswith(endpoint)) as response:
                        page.locator('#'+button).click()
                    result=response.value
                    if not result.ok:raise AssertionError(result.text())
                    expect(page.locator('.dw-panel.active fieldset')).to_be_enabled(timeout=15000)
                    return result.json()
                try:
                    page.goto(BASE,wait_until='domcontentloaded')
                    expect(page.locator('#productSplash')).to_be_hidden(timeout=20000)
                    expect(page.locator('#releaseDone')).to_be_visible();page.locator('#releaseDone').click()
                    nav('TopoSync','Point Clouds','pointcloudTools')
                    page.locator('#pcPath').fill(str(las))
                    inspected=call('pcInspect','/api/v9/pointcloud/inspect')
                    assert inspected['metadata']['point_count']==3, inspected
                    retained=call('pcImport','/api/v9/pointcloud/import')
                    assert retained['source']['sha256']==hashes[las],retained
                    expect(page.locator('#pcSources')).to_contain_text('sample.las')
                    cloud_status=api('/api/v9/pointcloud/status')
                    if cloud_status.get('laspy_ready'):
                        page.locator('#pcLimit').fill('2');preview=call('pcSample','/api/v9/pointcloud/sample')
                        assert preview['sample_count']==2,preview
                        report['optional']['LAS_sample']='PASS (real laspy)'
                    else:
                        expect(page.locator('#pcSample')).to_be_disabled()
                        report['optional']['LAS_sample']='NOT INSTALLED: metadata and retained source tested'
                    page.screenshot(path=str(OUT/'point-cloud-dark.png'))
                    report['checks'].append('G01 actual Inspect/Retain source with source SHA-256 and optional preview status')
                    nav('Home','Project Automation','workflowTools')
                    page.locator('#wfName').fill('Synthetic review workflow')
                    page.locator('#wfActions select').first.select_option('create_review_item')
                    page.locator('#wfActions textarea').first.fill('{"title":"Synthetic desktop review"}')
                    saved=call('wfSave','/api/v9/workflows')
                    run=call('wfRun','/api/v9/workflows/run');assert run['status']=='COMPLETED',run
                    expect(page.locator('#wfRunDetails')).to_contain_text('COMPLETED')
                    page.locator('#workflowTools summary').click()
                    export=call('wfExport','/api/v9/workflows/save-yaml')
                    text=Path(export['output_path']).read_text(encoding='utf-8');assert 'Synthetic review workflow' in text
                    page.locator('#wfYaml').fill(text)
                    assert call('wfImport','/api/v9/workflows/import-yaml')['count']==1
                    # Exercise a real gated action but reject it; no email or deliverable dispatch.
                    page.locator('#wfSelect').select_option(saved['workflow_id'])
                    page.locator('#wfActions select').first.select_option('send_notification')
                    page.locator('#wfActions textarea').first.fill('{}')
                    call('wfSave','/api/v9/workflows')
                    run=call('wfRun','/api/v9/workflows/run');assert run['status']=='WAITING_APPROVAL',run
                    expect(page.locator('#wfApprove')).to_be_enabled()
                    rejected=call('wfReject','/api/v9/workflows/approve');assert rejected['status']=='REJECTED'
                    expect(page.locator('#wfApprove')).to_be_disabled()
                    # No notification configuration exists in this synthetic project.
                    page.locator('#wfActions select').first.select_option('build_deliverable')
                    page.locator('#wfActions textarea').first.fill('{"label":"Synthetic approved delivery"}')
                    call('wfSave','/api/v9/workflows')
                    run=call('wfRun','/api/v9/workflows/run');assert run['status']=='WAITING_APPROVAL'
                    approved=call('wfApprove','/api/v9/workflows/approve');assert approved['status']=='COMPLETED',approved
                    page.screenshot(path=str(OUT/'workflow-dark.png'))
                    report['checks'].append('G02 actual editor/save/run/history/YAML export/import and explicit rejection/approval without contact')
                    nav('GISSync','CRS Diagnostics','crsTools')
                    page.locator('#cdSource').fill('EPSG:4326');page.locator('#cdTarget').fill('EPSG:3857')
                    page.locator('#cdX').fill('1');page.locator('#cdY').fill('0')
                    diag=call('cdOperations','/api/v9/crs/operations')
                    assert abs(diag['sample']['target']['x']-111319.49079327357)<1e-6,diag
                    page.screenshot(path=str(OUT/'crs-dark.png'))
                    report['checks'].append('G03 actual operation comparison independently checked 1 degree longitude')
                    nav('ReportSync','Excel Template Mapper','templateTools')
                    page.locator('#tmPath').fill(str(workbook));page.locator('#tmName').fill('Preserved company example')
                    call('tmInspect','/api/v9/reports/templates/inspect')
                    entry=call('tmRegister','/api/v9/reports/templates')
                    page.locator('#tmTableEnabled').check();page.locator('#tmSheet').select_option('Points')
                    for col,field in [('A','point_id'),('B','northing'),('C','easting'),('D','elevation')]:
                        page.locator('#tmAddColumn').click()
                        row=page.locator('#tmColumns .dw-mapping').last
                        row.locator('.tm-target').fill(col);row.locator('.tm-field').select_option(field)
                    call('tmSave','/api/v9/reports/templates/mapping')
                    rendered=call('tmRender','/api/v9/reports/templates/render')
                    out=Path(rendered['output_path']);assert out.is_file()
                    w=load_workbook(out)
                    try:
                        assert w['Points']['A2'].value=='001A' and w['Points']['A3'].value=='001B'
                        assert w['Cover']['B1'].value=='Remaining Workflow QA'
                        assert w['Cover']['D5'].value=='=1+1'
                    finally:w.close()
                    page.screenshot(path=str(OUT/'template-dark.png'))
                    report['checks'].append('G04 actual template inspection/retention/mapping/render; output reopened, IDs/formula preserved')
                    nav('GISSync','External GIS Processing','bridgeTools')
                    page.locator('#gbQgis').fill(str(temp/'missing_qgis.exe'));page.locator('#gbGrass').fill(str(temp/'missing_grass.exe'))
                    engines=call('gbDiscover','/api/v9/gis-bridges/status')
                    assert not engines['qgis']['ready'] and not engines['grass']['ready']
                    page.locator('#gbAlgorithm').fill('native:buffer')
                    page.locator('#gbAdd').click();row=page.locator('#gbParams .dw-mapping').last
                    row.locator('.gb-name').fill('DISTANCE');row.locator('.gb-type').select_option('number');row.locator('.gb-value').fill('1')
                    page.locator('#gbRun').click()
                    expect(page.locator('#bridgeTools .dw-error')).to_be_visible(timeout=15000)
                    page.screenshot(path=str(OUT/'gis-bridge-dark.png'))
                    report['checks'].append('G05 actual engine discovery/typed command controls; missing engine explicitly refused')
                    # Switch active project outside panel to ensure stale writes fail rather than hit the other job.
                    nav('Home','Project Automation','workflowTools')
                    api('/api/v9/project/create',{'parent_folder':str(temp/'projects'),'name':'Second QA project','crs':'EPSG:2278'})
                    page.locator('#wfName').fill('Must not leak into second project');page.locator('#wfSave').click()
                    expect(page.locator('#workflowTools .dw-error')).to_contain_text('Project changed',timeout=15000)
                    assert api('/api/v9/workflows')['workflows']==[]
                    report['checks'].append('Stale active panel blocked after project switch; second project unchanged')
                    api('/api/v9/config/ui',{'appearance':'light','theme':'classic','accent':'default'})
                    page.reload(wait_until='domcontentloaded');expect(page.locator('#productSplash')).to_be_hidden(timeout=15000)
                    for module,label,view in [('TopoSync','Point Clouds','pointcloudTools'),('Home','Project Automation','workflowTools'),('GISSync','CRS Diagnostics','crsTools'),('ReportSync','Excel Template Mapper','templateTools'),('GISSync','External GIS Processing','bridgeTools')]:
                        nav(module,label,view);expect(page.locator('html')).to_have_attribute('data-theme','light')
                        page.screenshot(path=str(OUT/(view+'-light.png')))
                        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
                    assert not errors,errors
                    assert all(hashlib.sha256(p.read_bytes()).hexdigest()==h for p,h in hashes.items())
                    report.update(status='PASS',js_errors=errors,originals_preserved=True,project_id=project_id)
                except Exception as exc:
                    report['failure']=f'{type(exc).__name__}: {exc}';report['js_errors']=errors
                    try:page.screenshot(path=str(OUT/'failure.png'))
                    except Exception:pass
                    raise
                finally:
                    context.tracing.stop(path=str(OUT/'trace.zip'));browser.close()
        finally:
            server.terminate()
            try:server.wait(timeout=10)
            except subprocess.TimeoutExpired:server.kill();server.wait()
            (OUT/'results.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report,indent=2))


if __name__=='__main__':main()
